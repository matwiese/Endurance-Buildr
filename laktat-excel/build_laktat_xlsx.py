#!/usr/bin/env python3
"""Erzeugt Laktatdiagnostik.xlsx – Excel-Nachbau der Webapp "Laktatdiagnostik – Auswertung".

Alle Ergebnisse sind Excel-Formeln (keine Makros), die sich bei jeder Eingabe neu berechnen.
Nichtlineare Suchen der Webapp (Nullstellen, Dmax, Breakpoint-Regressionen, Exponentialfit)
werden durch feine Rasterauswertungen mit Interpolation bzw. Verfeinerung nachgebildet.

Aufruf:  python3 build_laktat_xlsx.py [ziel.xlsx]
"""
import sys
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

OUT = sys.argv[1] if len(sys.argv) > 1 else "Laktatdiagnostik.xlsx"

# ------------------------------------------------------------------ Konstanten
MAXST = 20                    # max. Stufen in der Eingabetabelle
ST0 = 17                      # erste Datenzeile der Stufentabelle (Eingabe)
ST1 = ST0 + MAXST - 1
C0, C1 = 3, 3 + MAXST - 1     # kompakte Liste in Calc (Zeilen)
NG = 1000                     # Rasterintervalle im Messbereich
G0, G1 = 2, 2 + NG            # Grid-Zeilen

FONT = "Arial"
INK, MUTED, LT1C, LT2C, ACC = "18202B", "5B6676", "1F6FB2", "C8102E", "C8102E"
ZC = ["E3F1E5", "B9E3C0", "9CC7B2", "F6E49E", "F3B4A8"]
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=16, bold=True, color=INK)
f_h2 = Font(name=FONT, size=12, bold=True, color=INK)
f_muted = Font(name=FONT, size=9, color=MUTED)
f_input = Font(name=FONT, size=10, color="0000FF")
fill_in = PatternFill("solid", fgColor="FFF9D6")
fill_soft = PatternFill("solid", fgColor="EDF0F4")
fill_head = PatternFill("solid", fgColor="18202B")
thin = Side(style="thin", color="D8DEE6")
b_bottom = Border(bottom=thin)
b_box = Border(left=thin, right=thin, top=thin, bottom=thin)

METHODS = [
    ("ltp1h", "LTP1 (Hofmann/Pokan)", "LT1", "erster Laktatumstellpunkt"),
    ("ltp1", "LTP1 (3-Phasen global)", "LT1", "erster Laktatumstellpunkt"),
    ("loglog", "Log-Log (Beaver)", "LT1", "erster Laktatanstieg"),
    ("b05", "Baseline + 0,5 mmol/L", "LT1", "erster Laktatanstieg"),
    ("lemin", "Minimales Laktatäquivalent", "LT1", "erster Laktatanstieg"),
    ("fix2", "fix 2,0 mmol/L", "LT1", "niedriger fixer Anker"),
    ("ltp2h", "LTP2 (Hofmann/Pokan)", "LT2", "zweiter Laktatumstellpunkt"),
    ("ltp2", "LTP2 (3-Phasen global)", "LT2", "zweiter Laktatumstellpunkt"),
    ("dmax", "Dmax (Cheng)", "LT2", "LT2"),
    ("moddmax", "Modified Dmax (Bishop)", "LT2", "LT2"),
    ("lpdmax", "Log-Poly-ModDmax", "LT2", "LT2 / MLSS-Schätzer"),
    ("expdmax", "Exp-Dmax", "LT2", "LT2"),
    ("dickhuth", "Dickhuth (LEmin + 1,5)", "LT2", "individuelle obere Schwelle"),
    ("b10", "Baseline + 1,0 mmol/L", "LT2", "Übergangsbereich"),
    ("b15", "Baseline + 1,5 mmol/L", "LT2", "obere Schwelle (Simon-Typ)"),
    ("fix3", "fix 3,0 mmol/L", "FIX", "fixer Anker"),
    ("fix35", "fix 3,5 mmol/L", "FIX", "fixer Anker"),
    ("fix4", "fix 4,0 mmol/L (OBLA)", "FIX", "obere fixe Schwelle, ≠ MLSS"),
    ("baldari", "Baldari-Guidetti", "FIX", "MLSS-Näherung"),
    ("keul", "Keul (Tangente 51°)", "FIX", "historisch"),
    ("lmin", "Laktatminimum (Tegtbur)", "FIX", "nur bei LMT-Protokoll"),
]
MID = {m[0]: i for i, m in enumerate(METHODS)}
GRP_NAMES = {"LT1": "Erster Laktatanstieg (LT1)", "LT2": "Oberer Übergangsbereich (LT2)",
             "FIX": "Fixe, diskrete und historische Verfahren"}
MODELS = [("poly2", "Polynom 2. Grades"), ("poly3", "Polynom 3. Grades (Standard für Dmax)"),
          ("poly4", "Polynom 4. Grades"), ("exp", "Exponentiell  a + b·e^(c·x)")]
MODEL_SHORT = ["Polynom 2. Grades", "Polynom 3. Grades", "Polynom 4. Grades", "Exponentiell"]
EXAMPLES = {
    "Mayer Elisabeth · Laufband 22.09.2026": dict(
        meta=dict(name="Mayer Elisabeth", birth="15.07.2012", date="22.09.2026", height=155.1, mass=40.4,
                  tester="Ernst Johannes, BSc", stageInc="1,5 km/h", restLa=1.6, restGlu=81),
        stages=[["03:00", 6, 2.2, 145], ["06:00", 7.5, 2.3, 167], ["09:00", 9, 2.7, 184], ["12:00", 10.5, 3.5, 193],
                ["15:00", 12, 5.1, 200], ["18:00", 13.5, 8.2, 207], ["19:30", 14.25, 9.3, 208]]),
    "Lan Ryan · Laufband 22.09.2026": dict(
        meta=dict(name="Lan Ryan", birth="12.04.2012", date="22.09.2026", height=163.8, mass=47.5,
                  tester="Ernst Johannes, BSc", stageInc="1,5 km/h", restLa=1.1, restGlu=101),
        stages=[["03:00", 6, 1.6, 153], ["06:00", 7.5, 1.7, 164], ["09:00", 9, 2.1, 180], ["12:00", 10.5, 2.4, 190],
                ["15:00", 12, 3.1, 199], ["18:00", 13.5, 4.6, 211]]),
}

wb = Workbook()
NAMES = {}


def defname(name, sheet, ref):
    col = "".join(c for c in ref if c.isalpha())
    row = "".join(c for c in ref if c.isdigit())
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!${col}${row}")
    NAMES[name] = f"'{sheet}'!${col}${row}"


def style_ws(ws):
    ws.sheet_view.showGridLines = False


def put(ws, ref, value, font=None, fill=None, fmt=None, align=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font or f_base
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def arr(ws, ref, formula, **kw):
    """Matrixformel (Strg+Umschalt+Enter), funktioniert in allen Excel-Versionen und LibreOffice."""
    c = put(ws, ref, None, **kw)
    ws[ref] = ArrayFormula(ref, formula)
    return c


def FX(v, d=2):
    return f"FIXED({v},{d},TRUE)"


# Referenzen auf die kompakte, sortierte Stufenliste (Calc)
RK = f"Calc!$F${C0}:$F${C1}"     # laufende Nummer k
RX = f"Calc!$G${C0}:$G${C1}"     # Belastung
RY = f"Calc!$I${C0}:$I${C1}"     # Laktat
RM = f"Calc!$K${C0}:$K${C1}"     # Maske gültig (1/0)
GX = f"Grid!$B${G0}:$B${G1}"
GXa = f"Grid!$B${G0}:$B${G1-1}"
GXb = f"Grid!$B${G0+1}:$B${G1}"
GFCOL = ["D", "E", "F", "G"]       # f(x) je Modell
GDCOL = ["H", "I", "J", "K"]       # f'(x) je Modell
GLECOL = ["M", "N", "O", "P"]      # f(x)/x je Modell


def T(x):
    return f"(({x})-k_mx)/k_sx"


def fmodel(m, x):
    t = T(x)
    if m == 0:
        return f"(k_p2_0+{t}*(k_p2_1+{t}*k_p2_2))"
    if m == 1:
        return f"(k_p3_0+{t}*(k_p3_1+{t}*(k_p3_2+{t}*k_p3_3)))"
    if m == 2:
        return f"(k_p4_0+{t}*(k_p4_1+{t}*(k_p4_2+{t}*(k_p4_3+{t}*k_p4_4))))"
    return f"(k_ea+k_eb*EXP(k_ec*{t}))"


def fsel(x):
    t = T(x)
    return f"IF(k_mexp,k_ea+k_eb*EXP(k_ec*{t}),k_s0+{t}*(k_s1+{t}*(k_s2+{t}*(k_s3+{t}*k_s4))))"


def hrat(x):
    return f"IF(k_hrok,k_h0+k_h1*({x})+k_h2*({x})^2,\"\")"


def pace(x):
    return f"IF(N({x})>0,INT(ROUND(3600/({x}),0)/60)&\":\"&TEXT(MOD(ROUND(3600/({x}),0),60),\"00\"),\"–\")"


# ====================================================================== Blätter
ws_help = wb.active
ws_help.title = "Anleitung"
ws_in = wb.create_sheet("Eingabe")
ws_out = wb.create_sheet("Auswertung")
ws_meth = wb.create_sheet("Verfahren")
ws_dia = wb.create_sheet("Diagramme")
ws_rep = wb.create_sheet("Bericht")
ws_ex = wb.create_sheet("Beispiele")
ws_lst = wb.create_sheet("Listen")
ws_c = wb.create_sheet("Calc")
ws_g = wb.create_sheet("Grid")
ws_e = wb.create_sheet("ExpFit")
ws_h = wb.create_sheet("LTP")
ws_3 = wb.create_sheet("LTP3")
ws_gd = wb.create_sheet("Guide")
ws_v = wb.create_sheet("VCalc")
ws_cd = wb.create_sheet("ChartData")
for w in wb.worksheets:
    style_ws(w)

# ====================================================================== Listen
L = ws_lst
put(L, "A1", "id", f_bold); put(L, "B1", "Verfahren", f_bold); put(L, "C1", "Gruppe", f_bold); put(L, "D1", "Zielkonstrukt", f_bold)
for i, (mid, name, grp, ctr) in enumerate(METHODS):
    r = 2 + i
    put(L, f"A{r}", mid); put(L, f"B{r}", name); put(L, f"C{r}", grp); put(L, f"D{r}", ctr)
put(L, "F1", "Modell", f_bold)
for i, (mid, name) in enumerate(MODELS):
    put(L, f"F{2+i}", name); put(L, f"G{2+i}", mid)
LISTS = {
    "H": ["Laufband (km/h)", "Radergometer (W)"],
    "I": ["Kapillar Ohrläppchen", "Kapillar Finger", "venös"],
    "J": ["Vollblut", "Plasma", "Hämolysat"],
    "K": ["niedrigster Stufenwert", "Ruhelaktat"],
    "L": ["linear", "Polynom 2. Grades"],
    "M": ["korrigiert: nie unterhalb der Laktatkurve", "freie Regression durch die Messpunkte"],
    "N": ["Nein", "Ja"],
    "O": ["mg/dL", "mmol/L"],
}
for col, vals in LISTS.items():
    for i, v in enumerate(vals):
        put(L, f"{col}{2+i}", v)
ML = f"Listen!$B$2:$B${1+len(METHODS)}"
ws_lst.sheet_state = "hidden"


def dv_list(ws, ref, src):
    dv = DataValidation(type="list", formula1=src, allow_blank=False)
    dv.error = "Bitte einen Wert aus der Liste wählen."
    dv.errorTitle = "Ungültige Auswahl"
    ws.add_data_validation(dv)
    dv.add(ref)


# ====================================================================== Eingabe
E = ws_in
put(E, "A1", "Laktatdiagnostik – Eingabe", f_head)
put(E, "A2", "Gelb hinterlegte Felder mit blauer Schrift sind Eingaben. Alles andere wird berechnet. "
    "Dezimalzahlen wie in Excel üblich eingeben (z. B. 10,5). Leere Zeilen werden ignoriert, Stufen nach Belastung sortiert.",
    f_muted)
E.column_dimensions["A"].width = 17
E.column_dimensions["B"].width = 22
E.column_dimensions["C"].width = 13
E.column_dimensions["D"].width = 19
E.column_dimensions["E"].width = 20
E.column_dimensions["F"].width = 13
E.column_dimensions["G"].width = 3
E.column_dimensions["H"].width = 36
E.column_dimensions["I"].width = 36

put(E, "A3", "Athlet:in und Protokoll", f_h2)
meta_left = [("Name", "B4"), ("Geburtsdatum", "B5"), ("Testdatum", "B6"), ("Größe [cm]", "B7"), ("Masse [kg]", "B8"),
             ("Ergometer", "B9"), ("Testleitung", "B10"), ("Bemerkung", "B11"), ("Institution / Labor", "B12"),
             ("Adresse und Kontakt", "B13")]
meta_right = [("Stufendauer [min]", "E4"), ("Stufenhöhe", "E5"), ("Probenentnahme", "E6"), ("Matrix", "E7"),
              ("Analysator", "E8"), ("Ruhelaktat [mmol/L]", "E9"), ("Ruheglukose", "E10")]
for lab, ref in meta_left:
    put(E, f"A{ref[1:]}", lab, f_base)
    put(E, ref, None, f_input, fill_in, border=b_box)
for lab, ref in meta_right:
    put(E, f"D{ref[1:]}", lab, f_base)
    put(E, ref, None, f_input, fill_in, border=b_box)
E["D10"] = '="Ruheglukose ["&k_gluU&"]"'
for rr in (11, 12, 13):
    E.merge_cells(f"B{rr}:E{rr}")
    for cc in "CDE":
        E[f"{cc}{rr}"].fill = fill_in
E["B5"].number_format = "DD.MM.YYYY"
E["B6"].number_format = "DD.MM.YYYY"
E["E4"].number_format = "0.0#"
dv_list(E, "B9", "Listen!$H$2:$H$3")
dv_list(E, "E6", "Listen!$I$2:$I$4")
dv_list(E, "E7", "Listen!$J$2:$J$4")
E["B5"].comment = Comment("Datum als TT.MM.JJJJ eingeben (echtes Excel-Datum, für die Altersberechnung).", "Laktat")
E["E4"].comment = Comment("Stufendauer in Minuten als Zahl, z. B. 3 oder 2,5.", "Laktat")

put(E, "A15", "Messwerte je Stufe", f_h2)
heads = ["Stufe", "Zeit [mm:ss]", None, "Laktat [mmol/L]", "HF [1/min]", None]
for i, h in enumerate(heads):
    c = E.cell(16, 1 + i)
    c.value = h
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    c.fill = fill_head
    c.alignment = Alignment(horizontal="center", wrap_text=True)
E["C16"] = '="Belastung ["&k_unit&"]"'
E["F16"] = '="Glukose ["&k_gluU&"]"'
E.row_dimensions[16].height = 28
for i in range(MAXST):
    r = ST0 + i
    put(E, f"A{r}", i + 1, f_muted, align=Alignment(horizontal="center"))
    for cc in "BCDEF":
        put(E, f"{cc}{r}", None, f_input, fill_in, border=b_box)
    E[f"C{r}"].number_format = "0.00"
    E[f"D{r}"].number_format = "0.00"
    E[f"E{r}"].number_format = "0"
    E[f"F{r}"].number_format = "0.0"
put(E, f"A{ST1+2}", "Zeit als Text eingeben (z. B. 03:00). Glukose wird nur ausgewertet, wenn rechts „Glukose je Stufe erfassen = Ja“.", f_muted)

put(E, "H3", "Auswertung", f_h2)
settings = [
    (4, "Kurvenmodell Laktat", MODELS[1][1], "Listen!$F$2:$F$5"),
    (5, "LT1-Verfahren", "LTP1 (Hofmann/Pokan)", ML),
    (6, "LT2-Verfahren", "LTP2 (Hofmann/Pokan)", ML),
    (7, "LTP1-Suchbereich bis [% vmax]", 70, None),
    (8, "Baseline für Δ-Verfahren", "niedrigster Stufenwert", "Listen!$K$2:$K$3"),
    (9, "Modell Herzfrequenz", "Polynom 2. Grades", "Listen!$L$2:$L$3"),
    (10, "Hilfsgeraden-Konstruktion", LISTS["M"][0], "Listen!$M$2:$M$3"),
    (11, "Glukose je Stufe erfassen", "Nein", "Listen!$N$2:$N$3"),
    (12, "Glukose-Einheit", "mg/dL", "Listen!$O$2:$O$3"),
]
for r, lab, val, src in settings:
    put(E, f"H{r}", lab)
    put(E, f"I{r}", val, f_input, fill_in, border=b_box)
    if src:
        dv_list(E, f"I{r}", src)
dvr = DataValidation(type="decimal", operator="between", formula1="40", formula2="95")
dvr.error = "Wert zwischen 40 und 95 %."
E.add_data_validation(dvr)
dvr.add("I7")
put(E, "H14", "Schwellen manuell setzen (leer = Algorithmus)", f_h2)
put(E, "H15", None)
E["H15"] = '="LT1 manuell ["&k_unit&"]"'
E["H16"] = '="LT2 manuell ["&k_unit&"]"'
for r in (15, 16):
    put(E, f"I{r}", None, f_input, fill_in, border=b_box)
put(E, "H18", "Regeln für Trainingsbereiche", f_h2)
zr = [(19, "A1 ab [% LT1]", 81), (20, "A2 bis [% der Strecke LT1→LT2]", 52), (21, "A3 bis [% LT2]", 104.5),
      (22, "A4 bis [% vmax]", 98)]
for r, lab, v in zr:
    put(E, f"H{r}", lab)
    put(E, f"I{r}", v, f_input, fill_in, border=b_box)
put(E, "H23", "Voreinstellung (LSA-Beispielauswertung): 81 / 52 / 104,5 / 98. Konvention, keine physiologische Größe.", f_muted)
put(E, "H25", "Bereichsgrenzen manuell (leer = Regel)", f_h2)
edges = ["A0|A1", "A1|A2 (sonst = LT1)", "A2|A3", "A3|A4", "A4 Ende"]
for i, lab in enumerate(edges):
    E[f"H{26+i}"] = f'="{lab} ["&k_unit&"]"'
    E[f"H{26+i}"].font = f_base
    put(E, f"I{26+i}", None, f_input, fill_in, border=b_box)
put(E, "H32", "Beispieldaten: Blatt „Beispiele“. Zum Speichern eines Tests die Datei unter neuem Namen sichern.", f_muted)
E.freeze_panes = "A3"

# ====================================================================== Calc
C = ws_c
put(C, "A1", "Rechenblatt – kompakte Stufenliste, Parameter, Modellanpassung (nicht bearbeiten)", f_bold)
for j, h in enumerate(["Zeile", "Key x", "Key HF", "Key Glu", "", "k", "x", "Quelle", "La", "HF", "gültig", "t", "t²", "t³",
                       "t⁴", "ln x", "ln La", "fit P2", "fit P3", "fit P4", "fit Exp", "x (Diagr.)", "La (Diagr.)",
                       "La/x (Diagr.)", "ΔLa", "Baldari", "Guide", "", "k HF", "x HF", "x² HF", "HF", "x HF (D)",
                       "HF (D)", "", "k Glu", "x Glu", "Glu", "Δ", "Tangente"]):
    put(C, f"{CL(1+j)}2", h, f_bold)
for i in range(MAXST):
    r = C0 + i
    er = ST0 + i
    C[f"A{r}"] = i + 1
    C[f"B{r}"] = (f'=IF(AND(ISNUMBER(Eingabe!$C${er}),ISNUMBER(Eingabe!$D${er})),'
                  f'IF(AND(Eingabe!$C${er}>0,Eingabe!$D${er}>0),Eingabe!$C${er},""),"")')
    C[f"C{r}"] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(Eingabe!$E${er})),IF(Eingabe!$E${er}>0,B{r},""),"")'
    C[f"D{r}"] = f'=IF(AND(k_gluOn,ISNUMBER(B{r}),ISNUMBER(Eingabe!$F${er})),IF(Eingabe!$F${er}>0,B{r},""),"")'
    C[f"F{r}"] = i + 1
    C[f"G{r}"] = f"=IF(F{r}<=k_n,SMALL($B${C0}:$B${C1},F{r}),0)"
    C[f"H{r}"] = f"=IF(F{r}<=k_n,MATCH(G{r},$B${C0}:$B${C1},0),0)"
    C[f"I{r}"] = f"=IF(F{r}<=k_n,INDEX(Eingabe!$D${ST0}:$D${ST1},H{r}),0)"
    C[f"J{r}"] = f'=IF(F{r}<=k_n,IF(ISNUMBER(INDEX(Eingabe!$E${ST0}:$E${ST1},H{r})),INDEX(Eingabe!$E${ST0}:$E${ST1},H{r}),""),"")'
    C[f"K{r}"] = f"=IF(F{r}<=k_n,1,0)"
    C[f"L{r}"] = f"=IF(K{r}=1,(G{r}-k_mx)/k_sx,0)"
    C[f"M{r}"] = f"=L{r}^2"
    C[f"N{r}"] = f"=L{r}^3"
    C[f"O{r}"] = f"=L{r}^4"
    C[f"P{r}"] = f"=IF(K{r}=1,LN(G{r}),0)"
    C[f"Q{r}"] = f"=IF(K{r}=1,LN(I{r}),0)"
    C[f"R{r}"] = f"=IF(K{r}=1,{fmodel(0, f'G{r}')},0)"
    C[f"S{r}"] = f"=IF(K{r}=1,{fmodel(1, f'G{r}')},0)"
    C[f"T{r}"] = f"=IF(K{r}=1,{fmodel(2, f'G{r}')},0)"
    C[f"U{r}"] = f"=IF(K{r}=1,{fmodel(3, f'G{r}')},0)"
    C[f"V{r}"] = f"=IF(K{r}=1,G{r},NA())"
    C[f"W{r}"] = f"=IF(K{r}=1,I{r},NA())"
    C[f"X{r}"] = f"=IF(K{r}=1,I{r}/G{r},NA())"
    C[f"Y{r}"] = f"=IF(AND(F{r}>=2,F{r}<=k_n),ROUND(I{r}-I{r-1},9),0)" if i else "=0"
    nxt = f"Y{r+1}" if i < MAXST - 1 else "0"
    C[f"Z{r}"] = f"=IF(AND(F{r}>=2,F{r}<=k_n-1),IF(AND(Y{r}>=0.5,{nxt}>=0.5),1,0),0)"
    # Hilfsgeraden-Wert am Messpunkt (für RMSE)
    C[f"AA{r}"] = (f"=IF(K{r}=1,IF(G{r}<=k_gk1,k_gy0+(k_gy1-k_gy0)*(G{r}-k_gx0)/(k_gk1-k_gx0),"
                   f"IF(G{r}<=k_gk2,k_gy1+(k_gy2-k_gy1)*(G{r}-k_gk1)/(k_gk2-k_gk1),"
                   f"k_gy2+(k_gy3-k_gy2)*(G{r}-k_gk2)/(k_gx3-k_gk2))),0)")
    # HF-Liste
    C[f"AC{r}"] = i + 1
    C[f"AD{r}"] = f"=IF(AC{r}<=k_nh,SMALL($C${C0}:$C${C1},AC{r}),0)"
    C[f"AF{r}"] = f"=IF(AC{r}<=k_nh,INDEX(Eingabe!$E${ST0}:$E${ST1},MATCH(AD{r},$C${C0}:$C${C1},0)),0)"
    C[f"AE{r}"] = f"=AD{r}^2"
    C[f"AG{r}"] = f"=IF(AC{r}<=k_nh,AD{r},NA())"
    C[f"AH{r}"] = f"=IF(AC{r}<=k_nh,AF{r},NA())"
    # Glukose-Liste (monotone kubische Interpolation, wie Webapp)
    C[f"AJ{r}"] = i + 1
    C[f"AK{r}"] = f"=IF(AJ{r}<=k_ng,SMALL($D${C0}:$D${C1},AJ{r}),0)"
    C[f"AL{r}"] = f"=IF(AJ{r}<=k_ng,INDEX(Eingabe!$F${ST0}:$F${ST1},MATCH(AK{r},$D${C0}:$D${C1},0)),0)"
    C[f"AM{r}"] = f"=IF(AJ{r}<k_ng,(AL{r+1}-AL{r})/(AK{r+1}-AK{r}),0)"
    C[f"AR{r}"] = f"=IF(AJ{r}<=k_ng,AK{r},NA())"
    C[f"AS{r}"] = f"=IF(AJ{r}<=k_ng,AL{r},NA())"
# Tangenten der monotonen Interpolation (Fritsch-Carlson wie Webapp)
for i in range(MAXST):
    r = C0 + i
    prev = f"AM{r-1}" if i else "0"
    raw = (f"IF(AJ{r}=1,AM{r},IF(AJ{r}=k_ng,{prev},IF({prev}*AM{r}<=0,0,({prev}+AM{r})/2)))")
    C[f"AO{r}"] = f"=IF(AJ{r}<=k_ng,{raw},0)"
put(C, "AN2", "m roh", f_bold)
put(C, "AP2", "m links", f_bold)
put(C, "AQ2", "m rechts", f_bold)
# Begrenzung je Intervall (h>9 → Skalierung)
for i in range(MAXST):
    r = C0 + i
    C[f"AN{r}"] = f"=AO{r}"
    # links (m_i fürs Intervall i) und rechts (m_{i+1} fürs Intervall i)
    nx = f"AO{r+1}" if i < MAXST - 1 else "0"
    a = f"IF(AM{r}=0,0,AO{r}/AM{r})"
    b = f"IF(AM{r}=0,0,{nx}/AM{r})"
    hh = f"({a}^2+{b}^2)"
    C[f"AP{r}"] = f"=IF(AJ{r}<k_ng,IF(AM{r}=0,0,IF({hh}>9,3/SQRT({hh})*{a}*AM{r},AO{r})),0)"
    C[f"AQ{r}"] = f"=IF(AJ{r}<k_ng,IF(AM{r}=0,0,IF({hh}>9,3/SQRT({hh})*{b}*AM{r},{nx})),0)"

# ---------------- Parameter
P0 = 26
params = [
    ("k_n", "Anzahl gültiger Stufen", f"=COUNT($B${C0}:$B${C1})"),
    ("k_nh", "Anzahl HF-Werte", f"=COUNT($C${C0}:$C${C1})"),
    ("k_ng", "Anzahl Glukosewerte", f"=COUNT($D${C0}:$D${C1})"),
    ("k_unit", "Einheit", '=IF(Eingabe!$B$9="Radergometer (W)","W","km/h")'),
    ("k_run", "Laufband?", '=k_unit="km/h"'),
    ("k_xd", "Nachkommastellen Belastung", "=IF(k_run,2,0)"),
    ("k_gluOn", "Glukose erfassen", '=Eingabe!$I$11="Ja"'),
    ("k_gluU", "Glukose-Einheit", '=IF(Eingabe!$I$12="mmol/L","mmol/L","mg/dL")'),
    ("k_gd", "Glukose-Dezimalen", '=IF(k_gluU="mmol/L",1,0)'),
    ("k_dup", "doppelte Belastung", f"=SUMPRODUCT(($G${C0+1}:$G${C1}=$G${C0}:$G${C1-1})*($F${C0+1}:$F${C1}<=k_n))>0"),
    ("k_lo", "erste Stufe", f"=IF(k_n>0,$G${C0},0)"),
    ("k_hi", "letzte Stufe (vmax)", f"=IF(k_n>0,INDEX($G${C0}:$G${C1},k_n),1)"),
    ("k_rng", "Messbereich", "=MAX(k_hi-k_lo,1E-9)"),
    ("k_mx", "Mittelwert x", f"=IF(k_n>0,SUMPRODUCT({RM},{RX})/k_n,0)"),
    ("k_sx", "SD x (Population)", f"=IF(k_n>0,IF(SQRT(SUMPRODUCT({RM},({RX}-k_mx)^2)/k_n)>0,SQRT(SUMPRODUCT({RM},({RX}-k_mx)^2)/k_n),1),1)"),
    ("k_rest", "Ruhelaktat", '=IF(ISNUMBER(Eingabe!$E$9),Eingabe!$E$9,"")'),
    ("k_lamin", "min. Stufenlaktat", None),
    ("k_base", "Baseline", '=IF(AND(Eingabe!$I$8="Ruhelaktat",ISNUMBER(k_rest)),k_rest,k_lamin)'),
    ("k_lamax", "max. Laktat", f"=MAX({RY})"),
    ("k_ymean", "Mittel Laktat", f"=IF(k_n>0,SUMPRODUCT({RM},{RY})/k_n,0)"),
    ("k_sst", "SST Laktat", f"=SUMPRODUCT({RM},({RY}-k_ymean)^2)"),
    ("k_model", "Modell-Index", f'=IFERROR(MATCH(Eingabe!$I$4,Listen!$F$2:$F$5,0),2)'),
    ("k_mexp", "Modell exponentiell", "=k_model=4"),
    ("k_m1", "LT1-Verfahren (Index)", f"=IFERROR(MATCH(Eingabe!$I$5,{ML},0),1)"),
    ("k_m2", "LT2-Verfahren (Index)", f"=IFERROR(MATCH(Eingabe!$I$6,{ML},0),7)"),
    ("k_roi", "LTP1-Suchbereich (Anteil)", "=IF(AND(ISNUMBER(Eingabe!$I$7),Eingabe!$I$7>=40,Eingabe!$I$7<=95),Eingabe!$I$7/100,0.7)"),
    ("k_hrdeg", "HF-Modellgrad", '=IF(Eingabe!$I$9="linear",1,2)'),
    ("k_gfree", "Hilfsgeraden frei", '=Eingabe!$I$10="freie Regression durch die Messpunkte"'),
    ("k_x2nd", "zweite Stufe", f"=IF(k_n>1,$G${C0+1},k_lo)"),
    ("k_xn1", "vorletzte Stufe", f"=IF(k_n>1,INDEX($G${C0}:$G${C1},k_n-1),k_hi)"),
    ("k_yfirst", "Laktat erste Stufe", f"=$I${C0}"),
    ("k_ylast", "Laktat letzte Stufe", f"=IF(k_n>0,INDEX($I${C0}:$I${C1},k_n),0)"),
    ("k_dur", "Stufendauer [min]", '=IF(ISNUMBER(Eingabe!$E$4),Eingabe!$E$4,"")'),
]
for i, (nm, lab, f) in enumerate(params):
    r = P0 + i
    put(C, f"A{r}", lab)
    C[f"B{r}"] = f
    defname(nm, "Calc", f"B{r}")
    if nm == "k_lamin":
        arr(C, f"B{r}", f"=MIN(IF({RM}=1,{RY}))")
r = P0 + len(params)
# ---------------- Polynomfits (standardisiert wie in der Webapp)
PF = r + 1
put(C, f"A{PF}", "Polynom- und Exponentialfit (Koeffizienten in t = (x − x̄)/s)", f_bold)
r = PF + 1


def linest(deg, k):
    y = f"OFFSET($I${C0},0,0,k_n,1)"
    xs = f"OFFSET($L${C0},0,0,k_n,{deg})"
    return f"INDEX(LINEST({y},{xs}),1,{deg+1-k})"


for deg, pre, need in ((2, "p2", 3), (3, "p3", 4), (4, "p4", 5)):
    put(C, f"A{r}", f"Polynom {deg}. Grades verfügbar")
    C[f"B{r}"] = f"=AND(k_n>={need},NOT(k_dup))"
    defname(f"k_{pre}ok", "Calc", f"B{r}")
    r += 1
    for k in range(deg + 1):
        put(C, f"A{r}", f"{pre} c{k}")
        C[f"B{r}"] = f"=IF(k_{pre}ok,IFERROR({linest(deg, k)},0),0)"
        defname(f"k_{pre}_{k}", "Calc", f"B{r}")
        r += 1
put(C, f"A{r}", "Exponentialfit verfügbar"); C[f"B{r}"] = "=AND(k_n>=4,NOT(k_dup),ExpFit!$B$4<1E+299)"; defname("k_eok", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "exp a"); C[f"B{r}"] = "=IF(k_eok,ExpFit!$B$6,0)"; defname("k_ea", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "exp b (t)"); C[f"B{r}"] = "=IF(k_eok,ExpFit!$B$7,0)"; defname("k_eb", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "exp c (t)"); C[f"B{r}"] = "=IF(k_eok,ExpFit!$B$5,0)"; defname("k_ec", "Calc", f"B{r}"); r += 1
# Gütemaße je Modell
QF = r + 1
put(C, f"A{QF}", "Güte je Modell", f_bold)
for j, h in enumerate(["verfügbar", "Parameter", "SSE", "R²", "RMSE", "exakt"]):
    put(C, f"{CL(2+j)}{QF}", h, f_bold)
fitcols = ["R", "S", "T", "U"]
okn = ["k_p2ok", "k_p3ok", "k_p4ok", "k_eok"]
for m in range(4):
    rr = QF + 1 + m
    put(C, f"A{rr}", MODEL_SHORT[m])
    C[f"B{rr}"] = f"={okn[m]}"
    C[f"C{rr}"] = [3, 4, 5, 3][m]
    C[f"D{rr}"] = f"=SUMPRODUCT({RM},({RY}-Calc!${fitcols[m]}${C0}:${fitcols[m]}${C1})^2)"
    C[f"E{rr}"] = f'=IF(k_sst>0,1-D{rr}/k_sst,"")'
    C[f"F{rr}"] = f"=IF(k_n>0,SQRT(D{rr}/k_n),0)"
    C[f"G{rr}"] = f"=k_n<=C{rr}"
QT = f"Calc!$B${QF+1}:$G${QF+4}"
defname("k_fitok", "Calc", f"I{QF+1}")
C[f"H{QF+1}"] = "gewähltes Modell ok"
C[f"I{QF+1}"] = f"=INDEX($B${QF+1}:$B${QF+4},k_model)"
r = QF + 6
# gewählte Koeffizienten
put(C, f"A{r}", "Gewähltes Modell (Polynom-Koeffizienten, 0 bei exp)", f_bold); r += 1
for k in range(5):
    put(C, f"A{r}", f"s{k}")
    opts = [f"k_p2_{k}" if k <= 2 else "0", f"k_p3_{k}" if k <= 3 else "0", f"k_p4_{k}", "0"]
    C[f"B{r}"] = f"=CHOOSE(k_model,{','.join(opts)})"
    defname(f"k_s{k}", "Calc", f"B{r}")
    r += 1
put(C, f"A{r}", "R² gewählt"); C[f"B{r}"] = f"=INDEX($E${QF+1}:$E${QF+4},k_model)"; defname("k_r2", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "RMSE gewählt"); C[f"B{r}"] = f"=INDEX($F${QF+1}:$F${QF+4},k_model)"; defname("k_rmse", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "exakt gewählt"); C[f"B{r}"] = f"=INDEX($G${QF+1}:$G${QF+4},k_model)"; defname("k_exact", "Calc", f"B{r}"); r += 1

# Rohkoeffizienten des gewählten Modells (x statt t) für die Anzeige
put(C, f"A{r}", "Rohkoeffizienten a0..a4 (x in Belastungseinheit)", f_bold); r += 1
RAW0 = r
for j in range(5):
    terms = []
    for k in range(j, 5):
        from math import comb
        terms.append(f"k_s{k}/k_sx^{k}*{comb(k, j)}*(-k_mx)^{k-j}")
    put(C, f"A{r}", f"a{j}")
    C[f"B{r}"] = "=" + "+".join(terms)
    r += 1
put(C, f"A{r}", "Modellgleichung")
defname("k_eq", "Calc", f"B{r}")


def rawterm(j):
    v = f"$B${RAW0+j}"
    s = f"IF({v}=0,\"0\",FIXED(ABS({v}),MAX(0,4-INT(LOG10(ABS({v})))),TRUE))"
    sign = f'IF({v}<0," − "," + ")' if j else f'IF({v}<0,"−","")'
    return f'{sign}&{s}' + (f'&"·x{"" if j == 1 else ["", "", "²", "³", "⁴"][j]}"' if j else "")


maxdeg = "CHOOSE(k_model,2,3,4,0)"
poly_eq = '"La(x) = "&' + "&".join(
    [rawterm(0)] + [f'IF({maxdeg}>={j},{rawterm(j)},"")' for j in range(1, 5)])
def sig(v):
    return f'FIXED({v},MAX(0,4-INT(LOG10(ABS({v})+1E-300))),TRUE)'


exp_eq = ('"La(x) = "&' + sig("k_ea") + '&" + "&' + sig("k_eb*EXP(-k_ec*k_mx/k_sx)") + '&"·e^("&' + sig("k_ec/k_sx") + '&"·x)"')
C[f"B{r}"] = f"=IF(k_mexp,{exp_eq},{poly_eq})"
r += 2

# ---------------- HF-Modell
put(C, f"A{r}", "HF-Modell", f_bold); r += 1
HY = f"OFFSET($AF${C0},0,0,k_nh,1)"
put(C, f"A{r}", "HF-Modell verfügbar"); C[f"B{r}"] = "=k_nh>=3"; defname("k_hrok", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "Grad"); C[f"B{r}"] = "=MIN(k_hrdeg,k_nh-1)"; defname("k_hdeg", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "h0")
C[f"B{r}"] = (f"=IF(k_hrok,IF(k_hdeg=1,INDEX(LINEST({HY},OFFSET($AD${C0},0,0,k_nh,1)),1,2),"
              f"INDEX(LINEST({HY},OFFSET($AD${C0},0,0,k_nh,2)),1,3)),0)")
defname("k_h0", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "h1")
C[f"B{r}"] = (f"=IF(k_hrok,IF(k_hdeg=1,INDEX(LINEST({HY},OFFSET($AD${C0},0,0,k_nh,1)),1,1),"
              f"INDEX(LINEST({HY},OFFSET($AD${C0},0,0,k_nh,2)),1,2)),0)")
defname("k_h1", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "h2")
C[f"B{r}"] = f"=IF(AND(k_hrok,k_hdeg=2),INDEX(LINEST({HY},OFFSET($AD${C0},0,0,k_nh,2)),1,1),0)"
defname("k_h2", "Calc", f"B{r}"); r += 1
HRX = f"$AD${C0}:$AD${C1}"
HRY = f"$AF${C0}:$AF${C1}"
HRM = f"($AC${C0}:$AC${C1}<=k_nh)"
put(C, f"A{r}", "HF max"); C[f"B{r}"] = f'=IF(k_nh>0,MAX({HRY}),"")'; defname("k_hrmax", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "HF erste"); C[f"B{r}"] = f"=IF(k_nh>0,$AD${C0},0)"; defname("k_hx0", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "HF letzte x"); C[f"B{r}"] = f"=IF(k_nh>0,INDEX({HRX},k_nh),0)"; defname("k_hx1", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "HF letzte"); C[f"B{r}"] = f'=IF(k_nh>0,INDEX({HRY},k_nh),"")'; defname("k_hrlast", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "HF R²")
C[f"B{r}"] = (f'=IF(k_hrok,1-SUMPRODUCT({HRM}*({HRY}-(k_h0+k_h1*{HRX}+k_h2*{HRX}^2))^2)/'
              f'SUMPRODUCT({HRM}*({HRY}-SUMPRODUCT({HRM}*{HRY})/k_nh)^2),"")')
defname("k_hrr2", "Calc", f"B{r}"); r += 2

# ---------------- Log-Log (Beaver)
put(C, f"A{r}", "Log-Log (Beaver): Segmentgrenze k", f_bold); r += 1
LLH = r
for j, h in enumerate(["k", "gültig", "bL", "aL", "bR", "aR", "u*", "SSE", "Score"]):
    put(C, f"{CL(1+j)}{r}", h, f_bold)
r += 1
LL0 = r
U = f"$P${C0}"
V = f"$Q${C0}"
for k in range(2, MAXST - 1):
    rr = r
    C[f"A{rr}"] = k
    C[f"B{rr}"] = f"=AND(k_n>=4,A{rr}<=k_n-2)"
    uL = f"OFFSET({U},0,0,A{rr},1)"; vL = f"OFFSET({V},0,0,A{rr},1)"
    uR = f"OFFSET({U},A{rr},0,k_n-A{rr},1)"; vR = f"OFFSET({V},A{rr},0,k_n-A{rr},1)"
    C[f"C{rr}"] = f'=IF(B{rr},SLOPE({vL},{uL}),0)'
    C[f"D{rr}"] = f'=IF(B{rr},INTERCEPT({vL},{uL}),0)'
    C[f"E{rr}"] = f'=IF(B{rr},SLOPE({vR},{uR}),0)'
    C[f"F{rr}"] = f'=IF(B{rr},INTERCEPT({vR},{uR}),0)'
    C[f"G{rr}"] = f'=IF(AND(B{rr},C{rr}<>E{rr}),(F{rr}-D{rr})/(C{rr}-E{rr}),0)'
    C[f"H{rr}"] = f'=IF(B{rr},DEVSQ({vL})-C{rr}^2*DEVSQ({uL})+DEVSQ({vR})-E{rr}^2*DEVSQ({uR}),0)'
    C[f"I{rr}"] = (f'=IF(AND(B{rr},E{rr}>C{rr}),IF(AND(G{rr}>=$P${C0},G{rr}<=INDEX($P${C0}:$P${C1},k_n)),H{rr},1E+300),1E+300)')
    r += 1
LL1 = r - 1
put(C, f"A{r}", "Log-Log gefunden"); C[f"B{r}"] = f"=MIN($I${LL0}:$I${LL1})<1E+299"; defname("k_llok", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "Log-Log Zeile"); C[f"B{r}"] = f"=IF(k_llok,MATCH(MIN($I${LL0}:$I${LL1}),$I${LL0}:$I${LL1},0),1)"; defname("k_llrow", "Calc", f"B{r}"); r += 1
for nm, col in (("k_llk", "A"), ("k_llbl", "C"), ("k_llal", "D"), ("k_llbr", "E"), ("k_llar", "F"), ("k_llu", "G")):
    put(C, f"A{r}", nm); C[f"B{r}"] = f"=INDEX(${col}${LL0}:${col}${LL1},k_llrow)"; defname(nm, "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "Log-Log x"); C[f"B{r}"] = '=IF(k_llok,EXP(k_llu),"")'; defname("k_llx", "Calc", f"B{r}"); r += 2

# ---------------- ModDmax-Start, Baldari
put(C, f"A{r}", "Diskrete Verfahren", f_bold); r += 1
put(C, f"A{r}", "ModDmax: erste Stufe mit Anstieg > 0,4")
arr(C, f"B{r}", f"=IFERROR(MATCH(1,($Y${C0}:$Y${C1}>0.4)*($F${C0}:$F${C1}>=2)*($F${C0}:$F${C1}<=k_n),0),0)")
defname("k_mdk", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "ModDmax Start x"); C[f"B{r}"] = f"=IF(k_mdk>1,INDEX({RX},k_mdk-1),0)"; defname("k_mdx", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "ModDmax Start La"); C[f"B{r}"] = f"=IF(k_mdk>1,INDEX({RY},k_mdk-1),0)"; defname("k_mdy", "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "Baldari Stufe"); arr(C, f"B{r}", f"=IFERROR(MATCH(1,$Z${C0}:$Z${C1},0),0)"); defname("k_bdk", "Calc", f"B{r}"); r += 2

# ---------------- Hilfsgeraden (Knoten und Eckpunkte, befüllt aus Guide-Blatt)
put(C, f"A{r}", "Hilfsgeraden (aus Blatt Guide)", f_bold); r += 1
for nm, src in (("k_gok", "=Guide!$B$3"), ("k_gx0", "=k_lo"), ("k_gk1", "=Guide!$B$4"), ("k_gk2", "=Guide!$B$5"),
                ("k_gx3", "=k_hi"), ("k_gy0", "=Guide!$B$8"), ("k_gy1", "=Guide!$B$9"), ("k_gy2", "=Guide!$B$10"),
                ("k_gy3", "=Guide!$B$11")):
    put(C, f"A{r}", nm); C[f"B{r}"] = src; defname(nm, "Calc", f"B{r}"); r += 1
put(C, f"A{r}", "Hilfsgeraden RMSE")
C[f"B{r}"] = f'=IF(k_gok,SQRT(SUMPRODUCT({RM},({RY}-$AA${C0}:$AA${C1})^2)/k_n),"")'
defname("k_grmse", "Calc", f"B{r}"); r += 2
CALC_NEXT = r

# ====================================================================== ExpFit
X = ws_e
put(X, "A1", "Exponentialfit a + b·e^(c·t): Rastersuche über c mit zweifacher Verfeinerung (ersetzt Goldener-Schnitt-Suche)", f_bold)
Kx = f"Calc!$L${C0}:$L${C1}"
put(X, "A3", "Stufe 1 bestes c"); put(X, "A4", "min. SSE (final)"); put(X, "A5", "c final"); put(X, "A6", "a final"); put(X, "A7", "b final")


def exprow(ws, rr, cexpr):
    ws[f"A{rr}"] = cexpr
    e = f"EXP(A{rr}*{Kx})"
    ws[f"B{rr}"] = f"=SUMPRODUCT({RM},{e})"
    ws[f"C{rr}"] = f"=SUMPRODUCT({RM},EXP(2*A{rr}*{Kx}))"
    ws[f"D{rr}"] = f"=SUMPRODUCT({RM},{e},{RY})"
    ws[f"E{rr}"] = "=k_n*C{0}-B{0}^2".format(rr)
    ws[f"F{rr}"] = f"=IF(ABS(E{rr})<1E-14,0,(k_n*D{rr}-B{rr}*k_n*k_ymean)/E{rr})"
    ws[f"G{rr}"] = f"=IF(k_n>0,(k_n*k_ymean-F{rr}*B{rr})/k_n,0)"
    ws[f"H{rr}"] = f"=IF(AND(F{rr}>0,ABS(E{rr})>=1E-14),SUMPRODUCT({RM},({RY}-G{rr}-F{rr}*{e})^2),1E+300)"


for j, h in enumerate(["c", "Σe", "Σe²", "Σe·y", "det", "b", "a", "SSE"]):
    put(X, f"{CL(1+j)}9", h, f_bold)
E1a, E1b = 10, 409
for i in range(400):
    exprow(X, E1a + i, round(0.02 * (i + 1), 2))
X["B3"] = f"=INDEX(A{E1a}:A{E1b},MATCH(MIN(H{E1a}:H{E1b}),H{E1a}:H{E1b},0))"
E2a = E1b + 2
for i in range(41):
    exprow(X, E2a + i, f"=MAX(0.0001,$B$3-0.02)+($B$3+0.02-MAX(0.0001,$B$3-0.02))*{i}/40")
E2b = E2a + 40
X[f"J{E2a}"] = "bestes c Stufe 2"
X[f"K{E2a}"] = f"=INDEX(A{E2a}:A{E2b},MATCH(MIN(H{E2a}:H{E2b}),H{E2a}:H{E2b},0))"
E3a = E2b + 2
for i in range(41):
    exprow(X, E3a + i, f"=MAX(0.0001,$K${E2a}-0.001)+($K${E2a}+0.001-MAX(0.0001,$K${E2a}-0.001))*{i}/40")
E3b = E3a + 40
X["B4"] = f"=MIN(H{E3a}:H{E3b})"
X["B8"] = f"=MATCH(B4,H{E3a}:H{E3b},0)"
X["A8"] = "Zeile final"
X["B5"] = f"=INDEX(A{E3a}:A{E3b},B8)"
X["B6"] = f"=INDEX(G{E3a}:G{E3b},B8)"
X["B7"] = f"=INDEX(F{E3a}:F{E3b},B8)"
ws_e.sheet_state = "hidden"

# ====================================================================== Grid
G = ws_g
for j, h in enumerate(["i", "x", "t", "f P2", "f P3", "f P4", "f Exp", "f' P2", "f' P3", "f' P4", "f' Exp", "f'' P4",
                       "La/x P2", "La/x P3", "La/x P4", "La/x Exp", "f gewählt", "f' gewählt"]):
    put(G, f"{CL(1+j)}1", h, f_bold)
for i in range(NG + 1):
    r = G0 + i
    G[f"A{r}"] = i
    G[f"B{r}"] = f"=k_lo+k_rng*A{r}/{NG}"
    G[f"C{r}"] = f"=(B{r}-k_mx)/k_sx"
    t = f"C{r}"
    G[f"D{r}"] = f"=k_p2_0+{t}*(k_p2_1+{t}*k_p2_2)"
    G[f"E{r}"] = f"=k_p3_0+{t}*(k_p3_1+{t}*(k_p3_2+{t}*k_p3_3))"
    G[f"F{r}"] = f"=k_p4_0+{t}*(k_p4_1+{t}*(k_p4_2+{t}*(k_p4_3+{t}*k_p4_4)))"
    G[f"G{r}"] = f"=k_ea+k_eb*EXP(k_ec*{t})"
    G[f"H{r}"] = f"=(k_p2_1+2*k_p2_2*{t})/k_sx"
    G[f"I{r}"] = f"=(k_p3_1+{t}*(2*k_p3_2+3*k_p3_3*{t}))/k_sx"
    G[f"J{r}"] = f"=(k_p4_1+{t}*(2*k_p4_2+{t}*(3*k_p4_3+4*k_p4_4*{t})))/k_sx"
    G[f"K{r}"] = f"=k_eb*k_ec*EXP(k_ec*{t})/k_sx"
    G[f"L{r}"] = f"=(2*k_p4_2+{t}*(6*k_p4_3+12*k_p4_4*{t}))/k_sx^2"
    G[f"M{r}"] = f"=D{r}/B{r}"
    G[f"N{r}"] = f"=E{r}/B{r}"
    G[f"O{r}"] = f"=F{r}/B{r}"
    G[f"P{r}"] = f"=G{r}/B{r}"
    G[f"Q{r}"] = f"=CHOOSE(k_model,D{r},E{r},F{r},G{r})"
    G[f"R{r}"] = f"=CHOOSE(k_model,H{r},I{r},J{r},K{r})"
ws_g.sheet_state = "hidden"


def grng(col, a=G0, b=G1):
    return f"Grid!${col}${a}:${col}${b}"


# ====================================================================== LTP (Hofmann/Pokan)
H = ws_h
put(H, "A1", "LTP-Konzept nach Hofmann/Pokan: zwei sequenzielle Breakpoint-Regressionen (Raster 400 + Verfeinerung 40)", f_bold)


def seg2_block(ws, top, title, mask, lo, hi):
    """Zweisegment-Regression y = c0 + c1·x + c2·max(0, x−k) mit Knoten-Rastersuche. Liefert Zellbezüge."""
    put(ws, f"A{top}", title, f_bold)
    M = mask
    hdr = top + 1
    labels = ["nm", "Sx", "Sy", "Sxx", "Sxy", "Syy", "Cxx", "Cxy", "Cyy", "lo", "hi"]
    forms = [f"=SUMPRODUCT({M})", f"=SUMPRODUCT({M},{RX})", f"=SUMPRODUCT({M},{RY})",
             f"=SUMPRODUCT({M},{RX}^2)", f"=SUMPRODUCT({M},{RX},{RY})", f"=SUMPRODUCT({M},{RY}^2)",
             None, None, None, f"={lo}", f"={hi}"]
    ref = {}
    for j, (lb, fm) in enumerate(zip(labels, forms)):
        c = CL(2 + j)
        put(ws, f"{c}{hdr}", lb, f_bold)
        ref[lb] = f"${c}${hdr+1}"
        if fm:
            ws[f"{c}{hdr+1}"] = fm
    nm, Sx, Sy = ref["nm"], ref["Sx"], ref["Sy"]
    ws[ref["Cxx"].replace("$", "")] = f"=IF({nm}>0,{ref['Sxx']}-{Sx}^2/{nm},0)"
    ws[ref["Cxy"].replace("$", "")] = f"=IF({nm}>0,{ref['Sxy']}-{Sx}*{Sy}/{nm},0)"
    ws[ref["Cyy"].replace("$", "")] = f"=IF({nm}>0,{ref['Syy']}-{Sy}^2/{nm},0)"
    res = hdr + 3
    cols = ["k", "L", "gültig", "Sh", "Shh", "Sxh", "Shy", "Cxh", "Chh", "Chy", "det", "c2", "c1", "SSE", "Score"]

    def rows(r0, count, kexpr):
        for j, h in enumerate(cols):
            put(ws, f"{CL(1+j)}{r0-1}", h, f_bold)
        for i in range(count):
            rr = r0 + i
            ws[f"A{rr}"] = kexpr(i)
            hk = f"({RX}>A{rr})*({RX}-A{rr})"
            ws[f"B{rr}"] = f"=SUMPRODUCT({M}*({RX}<=A{rr}))"
            ws[f"C{rr}"] = f"=AND(B{rr}>=2,{nm}-B{rr}>=2)"
            ws[f"D{rr}"] = f"=SUMPRODUCT({M}*{hk})"
            ws[f"E{rr}"] = f"=SUMPRODUCT({M}*{hk}^2)"
            ws[f"F{rr}"] = f"=SUMPRODUCT({M}*{hk}*{RX})"
            ws[f"G{rr}"] = f"=SUMPRODUCT({M}*{hk}*{RY})"
            ws[f"H{rr}"] = f"=IF({nm}>0,F{rr}-{Sx}*D{rr}/{nm},0)"
            ws[f"I{rr}"] = f"=IF({nm}>0,E{rr}-D{rr}^2/{nm},0)"
            ws[f"J{rr}"] = f"=IF({nm}>0,G{rr}-D{rr}*{Sy}/{nm},0)"
            ws[f"K{rr}"] = f"={ref['Cxx']}*I{rr}-H{rr}^2"
            ws[f"L{rr}"] = f"=IF(ABS(K{rr})>1E-12,({ref['Cxx']}*J{rr}-H{rr}*{ref['Cxy']})/K{rr},0)"
            ws[f"M{rr}"] = f"=IF(ABS(K{rr})>1E-12,(I{rr}*{ref['Cxy']}-H{rr}*J{rr})/K{rr},0)"
            ws[f"N{rr}"] = f"={ref['Cyy']}-M{rr}*{ref['Cxy']}-L{rr}*J{rr}"
            ws[f"O{rr}"] = f"=IF(AND(C{rr},ABS(K{rr})>1E-12,L{rr}>0),N{rr},1E+300)"
        return r0, r0 + count - 1

    a0, a1 = rows(res + 1, 401, lambda i: f"={ref['lo']}+({ref['hi']}-{ref['lo']})*{i}/400")
    best = a1 + 2
    put(ws, f"A{best}", "bestes k (Raster)")
    ws[f"B{best}"] = f"=INDEX(A{a0}:A{a1},MATCH(MIN(O{a0}:O{a1}),O{a0}:O{a1},0))"
    ws[f"C{best}"] = f"=MIN(O{a0}:O{a1})"
    d = f"({ref['hi']}-{ref['lo']})/400"
    b0, b1 = rows(best + 3, 41, lambda i: f"=$B${best}-{d}+2*{d}*{i}/40")
    fin = b1 + 2
    out = {}
    for j, (lb, fm) in enumerate([
        ("gefunden", f"=MIN(O{b0}:O{b1},$C${best})<1E+299"),
        ("Zeile", f"=MATCH(MIN(O{b0}:O{b1}),O{b0}:O{b1},0)"),
        ("k", f"=IF(MIN(O{b0}:O{b1})<=$C${best},INDEX(A{b0}:A{b1},B{fin+1}),$B${best})"),
        ("s1", f"=INDEX(M{b0}:M{b1},B{fin+1})"),
        ("s2", f"=INDEX(M{b0}:M{b1},B{fin+1})+INDEX(L{b0}:L{b1},B{fin+1})"),
    ]):
        put(ws, f"A{fin+j}", lb)
        ws[f"B{fin+j}"] = fm
        out[lb] = f"LTP!$B${fin+j}"
    out["lo"] = f"LTP!{ref['lo']}"
    out["hi"] = f"LTP!{ref['hi']}"
    return out, fin + 6


put(H, "A3", "Parameter", f_bold)
H["A4"] = "Stufen im Suchbereich bis ROI"; H["B4"] = f"=SUMPRODUCT({RM}*({RX}<=k_roi*k_hi+1E-9))"
H["A5"] = "m1 (Stufen für LTP1)"; H["B5"] = "=MIN(k_n,MAX(B4,MIN(4,k_n)))"
H["A6"] = "Erweiterung LTP1"; H["B6"] = "=B5-B4"
m1mask = f"({RK}<=LTP!$B$5)*{RM}"
o1, nxt = seg2_block(H, 9, "LTP1-Suche (erste Stufe bis ROI)", m1mask, f"Calc!$G${C0}", f"INDEX({RX},LTP!$B$5)")
# LTP2-Teilmenge
H["D4"] = "Stufen < LTP1"; H["E4"] = f"=SUMPRODUCT({RM}*({RX}<{o1['k']}))"
H["D5"] = "Start s2"; H["E5"] = "=MIN(E4+1,MAX(1,k_n-3))"
H["D6"] = "Erweiterung LTP2"; H["E6"] = "=E4+1-E5"
m2mask = f"({RK}>=LTP!$E$5)*({RK}<=k_n)"
o2, _ = seg2_block(H, nxt, "LTP2-Suche (LTP1 bis vmax)", m2mask, f"INDEX({RX},LTP!$E$5)", "k_hi")
for nm, v in (("k_h1ok", f"=AND(k_n>=5,{o1['gefunden']})"), ("k_hk1", f"={o1['k']}"),
              ("k_h2ok", f"=AND(k_h1ok,{o2['gefunden']},{o2['k']}>{o1['k']})"), ("k_hk2", f"={o2['k']}")):
    rr = 4 + list(("k_h1ok", "k_hk1", "k_h2ok", "k_hk2")).index(nm)
    H[f"G{rr}"] = nm
    H[f"H{rr}"] = v
    defname(nm, "LTP", f"H{rr}")
HOF = dict(o1=o1, o2=o2)
ws_h.sheet_state = "hidden"

# ====================================================================== LTP3 (3-Phasen global)
Z = ws_3
put(Z, "A1", "3-Phasen-Regression (stetige stückweise Gerade mit zwei Knoten), Raster 61×61 + zwei Verfeinerungen 21×21", f_bold)
Z["A3"] = "Knotenbereich von"; Z["B3"] = "=k_x2nd"
Z["A4"] = "Knotenbereich bis"; Z["B4"] = "=k_xn1"
Z["A5"] = "Schritt grob"; Z["B5"] = "=MAX(B4-B3,1E-9)/60"
sums = [("Sx", f"=SUMPRODUCT({RM},{RX})"), ("Sy", f"=SUMPRODUCT({RM},{RY})"), ("Cxx", f"=SUMPRODUCT({RM},{RX}^2)-D3^2/k_n"),
        ("Cxy", f"=SUMPRODUCT({RM},{RX},{RY})-D3*D4/k_n"), ("Cyy", f"=SUMPRODUCT({RM},{RY}^2)-D4^2/k_n")]
for i, (lb, fm) in enumerate(sums):
    Z[f"C{3+i}"] = lb
    Z[f"D{3+i}"] = fm
Sx3, Sy3, Cxx3, Cxy3, Cyy3 = "LTP3!$D$3", "LTP3!$D$4", "LTP3!$D$5", "LTP3!$D$6", "LTP3!$D$7"
HC = ["k1", "k2", "a", "mid", "c", "ok", "S1", "S2", "S11", "S22", "S12", "Sx1", "Sx2", "S1y", "S2y",
      "C11", "C22", "C12", "Cx1", "Cx2", "C1y", "C2y", "det", "bx", "b1", "b2", "SSE", "Score"]


def hinge_formulas(k1, k2, c):
    """Formeln der Hinge-Regression für Knotenzellen k1,k2; c(name) liefert die Zelle eines Zwischenwerts."""
    h1 = f"({RX}>{k1})*({RX}-{k1})"
    h2 = f"({RX}>{k2})*({RX}-{k2})"
    F = {}
    F["a"] = f"=SUMPRODUCT({RM}*({RX}<={k1}))"
    F["mid"] = f"=SUMPRODUCT({RM}*({RX}>{k1})*({RX}<{k2}))"
    F["c"] = f"=SUMPRODUCT({RM}*({RX}>={k2}))"
    F["ok"] = f"=AND(k_n>=5,{k2}>{k1}+1E-9,{c('a')}>=2,{c('mid')}>=1,{c('c')}>=2)"
    F["S1"] = f"=SUMPRODUCT({RM}*{h1})"
    F["S2"] = f"=SUMPRODUCT({RM}*{h2})"
    F["S11"] = f"=SUMPRODUCT({RM}*{h1}^2)"
    F["S22"] = f"=SUMPRODUCT({RM}*{h2}^2)"
    F["S12"] = f"=SUMPRODUCT({RM}*{h1}*{h2})"
    F["Sx1"] = f"=SUMPRODUCT({RM}*{h1}*{RX})"
    F["Sx2"] = f"=SUMPRODUCT({RM}*{h2}*{RX})"
    F["S1y"] = f"=SUMPRODUCT({RM}*{h1}*{RY})"
    F["S2y"] = f"=SUMPRODUCT({RM}*{h2}*{RY})"
    n = "k_n"
    F["C11"] = f"={c('S11')}-{c('S1')}^2/{n}"
    F["C22"] = f"={c('S22')}-{c('S2')}^2/{n}"
    F["C12"] = f"={c('S12')}-{c('S1')}*{c('S2')}/{n}"
    F["Cx1"] = f"={c('Sx1')}-{Sx3}*{c('S1')}/{n}"
    F["Cx2"] = f"={c('Sx2')}-{Sx3}*{c('S2')}/{n}"
    F["C1y"] = f"={c('S1y')}-{c('S1')}*{Sy3}/{n}"
    F["C2y"] = f"={c('S2y')}-{c('S2')}*{Sy3}/{n}"
    a11, a12, a13 = Cxx3, c("Cx1"), c("Cx2")
    a22, a23, a33 = c("C11"), c("C12"), c("C22")
    b1, b2, b3 = Cxy3, c("C1y"), c("C2y")

    def det3(m):
        (p, q, r_), (s, t, u), (v, w, x) = m
        return f"(({p})*(({t})*({x})-({u})*({w}))-({q})*(({s})*({x})-({u})*({v}))+({r_})*(({s})*({w})-({t})*({v})))"
    F["det"] = "=" + det3([[a11, a12, a13], [a12, a22, a23], [a13, a23, a33]])
    dt = c("det")
    F["bx"] = f"=IF(ABS({dt})>1E-12,{det3([[b1, a12, a13], [b2, a22, a23], [b3, a23, a33]])}/{dt},0)"
    F["b1"] = f"=IF(ABS({dt})>1E-12,{det3([[a11, b1, a13], [a12, b2, a23], [a13, b3, a33]])}/{dt},0)"
    F["b2"] = f"=IF(ABS({dt})>1E-12,{det3([[a11, a12, b1], [a12, a22, b2], [a13, a23, b3]])}/{dt},0)"
    F["SSE"] = f"={Cyy3}-{c('bx')}*{Cxy3}-{c('b1')}*{c('C1y')}-{c('b2')}*{c('C2y')}"
    F["Score"] = f"=IF(AND({c('ok')},ABS({dt})>1E-12,{c('b1')}>0,{c('b2')}>0),{c('SSE')},1E+300)"
    return F


def hinge_table(ws, r0, count_side, k1expr, k2expr, title):
    put(ws, f"A{r0-2}", title, f_bold)
    for j, h in enumerate(HC):
        put(ws, f"{CL(1+j)}{r0-1}", h, f_bold)
    col = {h: CL(1 + j) for j, h in enumerate(HC)}
    rr = r0
    for i in range(count_side):
        for j in range(count_side):
            ws[f"A{rr}"] = k1expr(i)
            ws[f"B{rr}"] = k2expr(j)
            F = hinge_formulas(f"A{rr}", f"B{rr}", lambda nm, rr=rr: f"{col[nm]}{rr}")
            for nm, fm in F.items():
                ws[f"{col[nm]}{rr}"] = fm
            rr += 1
    return r0, rr - 1, col


T0 = 12
t0, t1, hc = hinge_table(Z, T0, 61, lambda i: f"=$B$3+$B$5*{i}", lambda j: f"=$B$3+$B$5*{j}", "Grobraster")
sc = hc["Score"]


def best_cells(ws, r, t0, t1, label):
    ws[f"F{r}"] = label
    ws[f"G{r}"] = f"=MIN({sc}{t0}:{sc}{t1})"
    ws[f"H{r}"] = f"=MATCH(G{r},{sc}{t0}:{sc}{t1},0)"
    ws[f"I{r}"] = f"=INDEX(A{t0}:A{t1},H{r})"
    ws[f"J{r}"] = f"=INDEX(B{t0}:B{t1},H{r})"


best_cells(Z, 3, t0, t1, "grob")
u0 = t1 + 4
u0a, u0b, _ = hinge_table(Z, u0, 21, lambda i: f"=$I$3-$B$5+$B$5*{i}/10", lambda j: f"=$J$3-$B$5+$B$5*{j}/10", "Verfeinerung 1")
best_cells(Z, 4, u0a, u0b, "fein 1")
v0 = u0b + 4
v0a, v0b, _ = hinge_table(Z, v0, 21, lambda i: f"=$I$4-$B$5/10+$B$5*{i}/100", lambda j: f"=$J$4-$B$5/10+$B$5*{j}/100", "Verfeinerung 2")
best_cells(Z, 5, v0a, v0b, "fein 2")
Z["F6"] = "final"
Z["G6"] = "=MIN(G3:G5)"
Z["H6"] = "=MATCH(G6,G3:G5,0)"
Z["I6"] = "=INDEX(I3:I5,H6)"
Z["J6"] = "=INDEX(J3:J5,H6)"
# Auswertung am Optimum (vertikaler Block)
put(Z, "AE2", "Auswertung am Optimum", f_bold)


def hinge_eval(ws, col_lab, col_val, r0, k1ref, k2ref, sheet):
    cells = {nm: f"{col_val}{r0+i}" for i, nm in enumerate(HC[2:])}
    F = hinge_formulas(k1ref, k2ref, lambda nm: cells[nm])
    for i, nm in enumerate(HC[2:]):
        ws[f"{col_lab}{r0+i}"] = nm
        ws[cells[nm]] = F[nm]
    extra = r0 + len(HC) - 2
    c = {nm: f"{sheet}!${col_val}${cells[nm][len(col_val):]}" for nm in cells}
    # Achsenabschnitt c0 (roh): f(x) = c0 + bx·x + b1·h1 + b2·h2
    ws[f"{col_lab}{extra}"] = "c0"
    ws[f"{col_val}{extra}"] = f"=({Sy3}-{cells['bx']}*{Sx3}-{cells['b1']}*{cells['S1']}-{cells['b2']}*{cells['S2']})/k_n"
    c["c0"] = f"{sheet}!${col_val}${extra}"
    return c, extra + 1


Z["AE3"] = "k1"; Z["AF3"] = "=I6"
Z["AE4"] = "k2"; Z["AF4"] = "=J6"
ev, nxt3 = hinge_eval(Z, "AE", "AF", 5, "$AF$3", "$AF$4", "LTP3")
Z[f"AE{nxt3}"] = "gefunden"; Z[f"AF{nxt3}"] = f"=AND(k_n>=5,G6<1E+299)"
defname("k_3ok", "LTP3", f"AF{nxt3}")
defname("k_3k1", "LTP3", "AF3")
defname("k_3k2", "LTP3", "AF4")
Z[f"AE{nxt3+1}"] = "RMSE"; Z[f"AF{nxt3+1}"] = f"=IF(k_3ok,SQRT(MAX(0,{ev['SSE']})/k_n),\"\")"
defname("k_3rmse", "LTP3", f"AF{nxt3+1}")
Z[f"AE{nxt3+2}"] = "f(k1)"; Z[f"AF{nxt3+2}"] = f"={ev['c0']}+{ev['bx']}*AF3"
defname("k_3f1", "LTP3", f"AF{nxt3+2}")
Z[f"AE{nxt3+3}"] = "f(k2)"; Z[f"AF{nxt3+3}"] = f"={ev['c0']}+{ev['bx']}*AF4+{ev['b1']}*(AF4-AF3)"
defname("k_3f2", "LTP3", f"AF{nxt3+3}")
Z[f"AE{nxt3+4}"] = "s1"; Z[f"AF{nxt3+4}"] = f"={ev['bx']}"; defname("k_3s1", "LTP3", f"AF{nxt3+4}")
Z[f"AE{nxt3+5}"] = "s2"; Z[f"AF{nxt3+5}"] = f"={ev['bx']}+{ev['b1']}"; defname("k_3s2", "LTP3", f"AF{nxt3+5}")
Z[f"AE{nxt3+6}"] = "s3"; Z[f"AF{nxt3+6}"] = f"={ev['bx']}+{ev['b1']}+{ev['b2']}"; defname("k_3s3", "LTP3", f"AF{nxt3+6}")
ws_3.sheet_state = "hidden"

# ====================================================================== VCalc (alle Verfahren × alle Modelle)
V = ws_v
put(V, "A1", "Alle Schwellenverfahren für alle vier Kurvenmodelle (Spalten je Modell: x, La, Fehler, Info, Warnung)", f_bold)
for m in range(4):
    c0 = 3 + 5 * m
    put(V, f"{CL(c0)}3", MODEL_SHORT[m], f_bold)
    for j, h in enumerate(["x", "La", "Fehler", "Info", "Warnung"]):
        put(V, f"{CL(c0+j)}4", h, f_bold)
MR0 = 5
for i, (mid, name, grp, ctr) in enumerate(METHODS):
    put(V, f"A{MR0+i}", mid)
    put(V, f"B{MR0+i}", name)
MATRIX = f"VCalc!$C${MR0}:$W${MR0+len(METHODS)-1}"

# Hilfsblöcke: Spalten C..F = Modelle, Zeilen ab 30
HB = 30
MC = ["C", "D", "E", "F"]
put(V, f"A{HB-1}", "Hilfsgrößen je Modell", f_bold)
for m in range(4):
    put(V, f"{MC[m]}{HB-1}", MODEL_SHORT[m], f_bold)
vr = HB


def vrow(label, fn, array=False):
    """Eine Zeile Hilfsgrößen: fn(m, col) liefert die Formel für Modell m."""
    global vr
    put(V, f"A{vr}", label)
    for m in range(4):
        f = fn(m, MC[m])
        if array:
            arr(V, f"{MC[m]}{vr}", f)
        else:
            V[f"{MC[m]}{vr}"] = f
    vr += 1
    return vr - 1


row_ok = vrow("Modell verfügbar", lambda m, c: f"={okn[m]}")
row_flo = vrow("f(erste Stufe)", lambda m, c: f"=Grid!${GFCOL[m]}${G0}")
row_fhi = vrow("f(letzte Stufe)", lambda m, c: f"=Grid!${GFCOL[m]}${G1}")


def level_block(label, Lexpr):
    """levelX der Webapp: erste steigende Lösung von f(x) = L im Messbereich."""
    rL = vrow(f"{label}: L", lambda m, c: f"={Lexpr(m, c)}")

    def FA(m):
        return grng(GFCOL[m], G0, G1 - 1)

    def FB(m):
        return grng(GFCOL[m], G0 + 1, G1)
    rc = vrow(f"{label}: Anzahl", lambda m, c: f"=SUMPRODUCT(({FA(m)}<{c}{rL})*({FB(m)}>={c}{rL}))")
    rp = vrow(f"{label}: Index", lambda m, c: f"=IF({c}{rc}>0,MATCH(1,({FA(m)}<{c}{rL})*({FB(m)}>={c}{rL}),0),0)", array=True)

    def xform(m, c):
        F = grng(GFCOL[m])
        fa, fb = f"INDEX({F},{c}{rp})", f"INDEX({F},{c}{rp}+1)"
        xa, xb = f"INDEX({GX},{c}{rp})", f"INDEX({GX},{c}{rp}+1)"
        return (f'=IF(NOT({c}${row_ok}),"",IF({c}{rL}=Grid!${GFCOL[m]}${G0},k_lo,IF({c}{rp}>0,'
                f'{xa}+({c}{rL}-{fa})/({fb}-{fa})*({xb}-{xa}),"")))')
    rx = vrow(f"{label}: x", xform)
    re_ = vrow(f"{label}: Fehler", lambda m, c: (
        f'=IF(NOT({c}${row_ok}),"Modell nicht anpassbar",IF({c}{rL}<{c}${row_flo},{FX(f"{c}{rL}")}&" mmol/L liegt unter der ersten Stufe",'
        f'IF({c}{rL}>{c}${row_fhi},{FX(f"{c}{rL}")}&" mmol/L wird im Test nicht erreicht",IF(ISNUMBER({c}{rx}),"","keine steigende Lösung"))))'))
    rw = vrow(f"{label}: Warnung", lambda m, c: f'=IF(AND(ISNUMBER({c}{rx}),{c}{rc}>1),"mehrere Lösungen, Fit nicht monoton","")')
    return dict(L=rL, x=rx, err=re_, warn=rw)


def dmax_block(label, xS, yS, xE, yE, model_override=None):
    """Dmax: größter Abstand der Kurve unterhalb der Sehne (Raster + Parabelverfeinerung)."""
    def mm(m):
        return model_override if model_override is not None else m
    rxs = vrow(f"{label}: xS", lambda m, c: f"={xS(m, c)}")
    rys = vrow(f"{label}: yS", lambda m, c: f"={yS(m, c)}")
    rxe = vrow(f"{label}: xE", lambda m, c: f"={xE(m, c)}")
    rye = vrow(f"{label}: yE", lambda m, c: f"={yE(m, c)}")
    rA = vrow(f"{label}: A", lambda m, c: f"={c}{rye}-{c}{rys}")
    rB = vrow(f"{label}: B", lambda m, c: f"=-({c}{rxe}-{c}{rxs})")
    rC = vrow(f"{label}: C", lambda m, c: f"={c}{rxe}*{c}{rys}-{c}{rye}*{c}{rxs}")
    rD = vrow(f"{label}: Norm", lambda m, c: f"=MAX(SQRT({c}{rA}^2+{c}{rB}^2),1E-12)")

    def dist(m, c, xr, fr):
        return f"({c}${rA}*{xr}+{c}${rB}*{fr}+{c}${rC})/{c}${rD}"

    def mask(c):
        return f"({GX}>{c}{rxs})*({GX}<{c}{rxe})"
    F = lambda m: grng(GFCOL[mm(m)])
    rmx = vrow(f"{label}: max. Abstand", lambda m, c: f"=MAX(IF({mask(c)},{dist(m, c, GX, F(m))},-1E+300))", array=True)
    rp = vrow(f"{label}: Index", lambda m, c: f"=IFERROR(MATCH({c}{rmx},IF({mask(c)},{dist(m, c, GX, F(m))},-1E+300),0),0)", array=True)

    def xform(m, c):
        p = f"{c}{rp}"
        dm = dist(m, c, f"INDEX({GX},{p}-1)", f"INDEX({F(m)},{p}-1)")
        d0 = dist(m, c, f"INDEX({GX},{p})", f"INDEX({F(m)},{p})")
        dp = dist(m, c, f"INDEX({GX},{p}+1)", f"INDEX({F(m)},{p}+1)")
        den = f"({dm}-2*{d0}+{dp})"
        okc = f"{c}${row_ok}" if model_override is None else f"$F${row_ok}"
        return (f'=IF(OR(NOT({okc}),{c}{rxe}<={c}{rxs},{c}{rmx}<=0,{p}<2,{p}>{NG}),"",'
                f'IF(OR(INDEX({GX},{p}-1)<={c}{rxs},INDEX({GX},{p}+1)>={c}{rxe}),"",'
                f'INDEX({GX},{p})+IF({den}<0,0.5*({dm}-{dp})/{den},0)*k_rng/{NG}))')
    rx = vrow(f"{label}: x", xform)
    rla = vrow(f"{label}: La", lambda m, c: f'=IF(ISNUMBER({c}{rx}),{fmodel(mm(m), f"{c}{rx}")},"")')
    return dict(x=rx, la=rla, xs=rxs, ys=rys, xe=rxe, ye=rye)


def root_block(label, cond, dcol_of):
    """erste Nullstelle einer Rasterbedingung (Vorzeichenwechsel) mit linearer Interpolation."""
    rp = vrow(f"{label}: Index", lambda m, c: f"=IFERROR(MATCH(1,{cond(m, c)},0),0)", array=True)

    def xform(m, c):
        p = f"{c}{rp}"
        D, s = dcol_of(m)
        da, db = f"(INDEX({D},{p})-({s}))", f"(INDEX({D},{p}+1)-({s}))"
        xa, xb = f"INDEX({GX},{p})", f"INDEX({GX},{p}+1)"
        return f'=IF(OR(NOT({c}${row_ok}),{p}=0),"",IF({db}={da},{xa},{xa}+(0-{da})/({db}-{da})*({xb}-{xa})))'
    rx = vrow(f"{label}: x", xform)
    rla = vrow(f"{label}: La", lambda m, c: f'=IF(ISNUMBER({c}{rx}),{fmodel(m, f"{c}{rx}")},"")')
    return dict(x=rx, la=rla)


# Laktatäquivalent-Minimum je Modell
def lemin_rows():
    LE = lambda m: grng(GLECOL[m])
    rv = vrow("LEmin: Minimum", lambda m, c: f"=MIN({LE(m)})")
    rp = vrow("LEmin: Index", lambda m, c: f"=MATCH({c}{rv},{LE(m)},0)")

    def xform(m, c):
        p = f"{c}{rp}"
        lm, l0, lp = f"INDEX({LE(m)},{p}-1)", f"INDEX({LE(m)},{p})", f"INDEX({LE(m)},{p}+1)"
        den = f"({lm}-2*{l0}+{lp})"
        return (f'=IF(NOT({c}${row_ok}),"",INDEX({GX},{p})+IF(AND({p}>1,{p}<{NG+1}),IF({den}>0,0.5*({lm}-{lp})/{den},0),0)*k_rng/{NG})')
    rx = vrow("LEmin: x", xform)
    rla = vrow("LEmin: La", lambda m, c: f'=IF(ISNUMBER({c}{rx}),{fmodel(m, f"{c}{rx}")},"")')
    rb = vrow("LEmin: am Rand", lambda m, c: f'=IF(ISNUMBER({c}{rx}),OR({c}{rx}-k_lo<k_rng*0.02,k_hi-{c}{rx}<k_rng*0.02),FALSE)')
    return dict(x=rx, la=rla, edge=rb)


LEM = lemin_rows()
LV = {}
for key, lab, expr in (("b05", "Baseline+0,5", lambda m, c: "k_base+0.5"), ("fix2", "fix 2", lambda m, c: "2"),
                       ("dickhuth", "Dickhuth", lambda m, c: f'IF(ISNUMBER({c}${LEM["la"]}),{c}${LEM["la"]}+1.5,1E+300)'),
                       ("b10", "Baseline+1,0", lambda m, c: "k_base+1"), ("b15", "Baseline+1,5", lambda m, c: "k_base+1.5"),
                       ("fix3", "fix 3", lambda m, c: "3"), ("fix35", "fix 3,5", lambda m, c: "3.5"),
                       ("fix4", "fix 4", lambda m, c: "4")):
    LV[key] = level_block(lab, expr)
lastx = f"INDEX({RX},k_n)"
lasty = "k_ylast"
DM = dmax_block("Dmax", lambda m, c: "k_lo", lambda m, c: "k_yfirst", lambda m, c: lastx, lambda m, c: lasty)
MD = dmax_block("ModDmax", lambda m, c: "k_mdx", lambda m, c: "k_mdy", lambda m, c: lastx, lambda m, c: lasty)
LPD = dmax_block("Log-Poly-ModDmax", lambda m, c: 'IF(k_llok,k_llx,0)',
                 lambda m, c: f'IF(k_llok,{fmodel(m, "k_llx")},0)', lambda m, c: lastx, lambda m, c: lasty)
EXD = dmax_block("Exp-Dmax", lambda m, c: "k_lo", lambda m, c: "k_yfirst", lambda m, c: lastx, lambda m, c: lasty,
                 model_override=3)
TAN51 = "TAN(RADIANS(51))"
KEUL = root_block("Keul", lambda m, c: (f"(({grng(GDCOL[m], G0, G1-1)}-{TAN51})*({grng(GDCOL[m], G0+1, G1)}-{TAN51})<=0)*1"),
                  lambda m: (grng(GDCOL[m]), TAN51))
LMIN = root_block("Laktatminimum", lambda m, c: (f"({grng(GDCOL[m], G0, G1-1)}<0)*({grng(GDCOL[m], G0+1, G1)}>=0)*"
                                                 f"({GXa}>k_lo+0.02*k_rng)*({GXa}<k_hi-0.02*k_rng)"),
                  lambda m: (grng(GDCOL[m]), "0"))
# Hinweise: min. Steigung und Wendepunkte (gewähltes Modell)
put(V, f"A{vr+1}", "gewähltes Modell: min f' ab 2. Stufe")
arr(V, f"C{vr+1}", f"=MIN(IF({GX}>=k_x2nd,{grng('R')}))")
defname("k_mind1", "VCalc", f"C{vr+1}")
put(V, f"A{vr+2}", "Poly4: Anzahl Wendepunkte")
V[f"C{vr+2}"] = f"=SUMPRODUCT(({grng('L', G0, G1-1)}*{grng('L', G0+1, G1)}<0)*1)"
defname("k_nwp", "VCalc", f"C{vr+2}")
vr += 4


def mat(m, i, j):
    """Zelle der Ergebnis-Matrix für Modell m, Verfahren i, Feld j (0=x,1=La,2=Fehler,3=Info,4=Warnung)."""
    return f"{CL(3 + 5 * m + j)}{MR0 + i}"


def set_method(mid, fx, fla, ferr, finfo, fwarn):
    i = MID[mid]
    for m in range(4):
        c = MC[m]
        vals = [fx(m, c), fla(m, c), ferr(m, c), finfo(m, c), fwarn(m, c)]
        for j, v in enumerate(vals):
            V[mat(m, i, j)] = "=" + v


E_ = lambda m, c: '""'
nok = lambda c: f'NOT({c}${row_ok})'
# LTP1/LTP2 Hofmann
o1, o2 = HOF["o1"], HOF["o2"]
set_method("ltp1h",
           lambda m, c: f'IF(AND(k_h1ok,{c}${row_ok}),k_hk1,"")',
           lambda m, c: f'IF(AND(k_h1ok,{c}${row_ok}),{fmodel(m, "k_hk1")},"")',
           lambda m, c: f'IF({nok(c)},"Modell nicht anpassbar",IF(k_n<5,"mind. 5 Stufen nötig",IF(k_h1ok,"","kein Anstieg im LTP1-Suchbereich")))',
           lambda m, c: f'IF(k_h1ok,"Suchbereich "&{FX(o1["lo"])}&"–"&{FX(o1["hi"])}&", Steigung "&{FX(o1["s1"])}&" → "&{FX(o1["s2"])},"")',
           lambda m, c: f'IF(AND(k_h1ok,LTP!$B$6>0),"Suchbereich um "&LTP!$B$6&" Stufe"&IF(LTP!$B$6>1,"n","")&" erweitert (zu wenige Stufen bis "&ROUND(k_roi*100,0)&" % vmax)","")')
set_method("ltp2h",
           lambda m, c: f'IF(AND(k_h2ok,{c}${row_ok}),k_hk2,"")',
           lambda m, c: f'IF(AND(k_h2ok,{c}${row_ok}),{fmodel(m, "k_hk2")},"")',
           lambda m, c: f'IF({nok(c)},"Modell nicht anpassbar",IF(k_n<5,"mind. 5 Stufen nötig",IF(NOT(k_h1ok),"kein Anstieg im LTP1-Suchbereich",IF(k_h2ok,"","kein zweiter Knick oberhalb LTP1"))))',
           lambda m, c: f'IF(k_h2ok,"Suchbereich LTP1–vmax, Steigung "&{FX(o2["s1"])}&" → "&{FX(o2["s2"])},"")',
           lambda m, c: f'IF(AND(k_h2ok,LTP!$E$6>0),"Suchbereich um "&LTP!$E$6&" Stufe"&IF(LTP!$E$6>1,"n","")&" unter LTP1 erweitert","")')
set_method("ltp1",
           lambda m, c: 'IF(k_3ok,k_3k1,"")', lambda m, c: 'IF(k_3ok,k_3f1,"")',
           lambda m, c: 'IF(k_3ok,"",IF(k_n<5,"mind. 5 Stufen nötig","keine konvexe 3-Phasen-Lösung"))',
           lambda m, c: f'IF(k_3ok,"Knick der Hilfsgeraden, RMSE "&{FX("k_3rmse", 3)},"")',
           lambda m, c: 'IF(AND(k_3ok,k_n<7),"wenige Stufen für 3 Geraden","")')
set_method("ltp2",
           lambda m, c: 'IF(k_3ok,k_3k2,"")', lambda m, c: 'IF(k_3ok,k_3f2,"")',
           lambda m, c: 'IF(k_3ok,"",IF(k_n<5,"mind. 5 Stufen nötig","keine konvexe 3-Phasen-Lösung"))',
           lambda m, c: f'IF(k_3ok,"Steigungen "&{FX("k_3s1")}&" / "&{FX("k_3s2")}&" / "&{FX("k_3s3")},"")',
           lambda m, c: 'IF(AND(k_3ok,k_n<7),"wenige Stufen für 3 Geraden","")')
set_method("loglog",
           lambda m, c: f'IF(AND(k_llok,{c}${row_ok}),k_llx,"")',
           lambda m, c: f'IF(AND(k_llok,{c}${row_ok}),{fmodel(m, "k_llx")},"")',
           lambda m, c: f'IF({nok(c)},"Modell nicht anpassbar",IF(k_llok,"","kein gültiger Breakpoint"))',
           lambda m, c: 'IF(k_llok,"Segmente "&k_llk&" + "&(k_n-k_llk)&" Punkte","")',
           lambda m, c: 'IF(AND(k_llok,k_n<8),"wenige Punkte je Segment, Breakpoint springt leicht","")')
for key, info in (("b05", lambda m, c: f'IF(ISNUMBER({c}{LV["b05"]["x"]}),"Baseline "&{FX("k_base")},"")'),
                  ("fix2", E_), ("b10", E_), ("b15", E_), ("fix3", E_), ("fix35", E_), ("fix4", E_),
                  ("dickhuth", lambda m, c: 'IF(ISNUMBER(' + c + str(LV["dickhuth"]["x"]) + '),"LEmin "&' + FX(c + str(LEM["x"])) + '&" bei "&' + FX(c + str(LEM["la"])) + '&" mmol/L","")')):
    lv = LV[key]
    warn = (lambda m, c, lv=lv: f'IF(AND(ISNUMBER({c}{lv["x"]}),{c}{LEM["edge"]}),"LEmin am Rand des Messbereichs",{c}{lv["warn"]})') \
        if key == "dickhuth" else (lambda m, c, lv=lv: f'{c}{lv["warn"]}')
    set_method(key, lambda m, c, lv=lv: f'{c}{lv["x"]}',
               lambda m, c, lv=lv: f'IF(ISNUMBER({c}{lv["x"]}),{c}{lv["L"]},"")',
               lambda m, c, lv=lv: f'{c}{lv["err"]}', info, warn)
set_method("lemin", lambda m, c: f'{c}{LEM["x"]}', lambda m, c: f'{c}{LEM["la"]}',
           lambda m, c: f'IF({nok(c)},"Modell nicht anpassbar","")',
           lambda m, c: f'IF(ISNUMBER({c}{LEM["x"]}),"min La/x","")',
           lambda m, c: f'IF({c}{LEM["edge"]},"Minimum am Rand des Messbereichs","")')


def dm_err(b):
    return lambda m, c: f'IF({nok(c)},"Modell nicht anpassbar",IF(ISNUMBER({c}{b["x"]}),"","kein Distanzmaximum"))'


set_method("dmax", lambda m, c: f'{c}{DM["x"]}', lambda m, c: f'{c}{DM["la"]}', dm_err(DM),
           lambda m, c: f'IF(ISNUMBER({c}{DM["x"]}),"Sehne erster → letzter Punkt","")', E_)
set_method("moddmax", lambda m, c: f'IF(k_mdk>1,{c}{MD["x"]},"")', lambda m, c: f'IF(k_mdk>1,{c}{MD["la"]},"")',
           lambda m, c: f'IF({nok(c)},"Modell nicht anpassbar",IF(k_mdk<2,"kein Anstieg > 0,4 mmol/L",IF(ISNUMBER({c}{MD["x"]}),"","kein Distanzmaximum")))',
           lambda m, c: f'IF(AND(k_mdk>1,ISNUMBER({c}{MD["x"]})),"Start bei "&{FX("k_mdx")},"")',
           lambda m, c: f'IF(AND(k_mdk>1,ISNUMBER({c}{MD["x"]}),k_n-k_mdk+1<3),"nur wenige Punkte nach dem Startpunkt","")')
set_method("lpdmax", lambda m, c: f'IF(k_llok,{c}{LPD["x"]},"")', lambda m, c: f'IF(k_llok,{c}{LPD["la"]},"")',
           lambda m, c: f'IF({nok(c)},"Modell nicht anpassbar",IF(NOT(k_llok),"Log-Log nicht bestimmbar",IF(ISNUMBER({c}{LPD["x"]}),"","kein Distanzmaximum")))',
           lambda m, c: f'IF(AND(k_llok,ISNUMBER({c}{LPD["x"]})),"Start bei Log-Log "&{FX("k_llx")},"")', E_)
set_method("expdmax", lambda m, c: f'IF(k_eok,{c}{EXD["x"]},"")', lambda m, c: f'IF(k_eok,{c}{EXD["la"]},"")',
           lambda m, c: f'IF(NOT(k_eok),"Exponentialfit nicht möglich",IF(ISNUMBER({c}{EXD["x"]}),"","kein Distanzmaximum"))',
           lambda m, c: f'IF(ISNUMBER({c}{EXD["x"]}),"eigener Fit a+b·e^(cx)","")', E_)
set_method("baldari", lambda m, c: f'IF(k_bdk>0,INDEX({RX},k_bdk),"")', lambda m, c: f'IF(k_bdk>0,INDEX({RY},k_bdk),"")',
           lambda m, c: 'IF(k_bdk>0,"","keine zwei aufeinanderfolgenden Anstiege ≥ 0,5")',
           lambda m, c: 'IF(k_bdk>0,"diskret, Messwert der Stufe","")',
           lambda m, c: 'IF(k_bdk>0,"abhängig von Stufenhöhe und Messrauschen","")')
set_method("keul", lambda m, c: f'IF(k_run,{c}{KEUL["x"]},"")', lambda m, c: f'IF(k_run,{c}{KEUL["la"]},"")',
           lambda m, c: f'IF({nok(c)},"Modell nicht anpassbar",IF(NOT(k_run),"nur für km/h definiert",IF(ISNUMBER({c}{KEUL["x"]}),"","Steigung 1,23 nicht erreicht")))',
           E_, lambda m, c: f'IF(AND(k_run,ISNUMBER({c}{KEUL["x"]})),"einheiten- und skalierungsabhängig","")')
set_method("lmin", lambda m, c: f'{c}{LMIN["x"]}', lambda m, c: f'{c}{LMIN["la"]}',
           lambda m, c: f'IF({nok(c)},"Modell nicht anpassbar",IF(ISNUMBER({c}{LMIN["x"]}),"","kein inneres Minimum (kein LMT-Protokoll)"))',
           E_, lambda m, c: f'IF(ISNUMBER({c}{LMIN["x"]}),"nur nach Vorbelastung interpretierbar","")')
# kompakter Hinweistext je Verfahren/Modell (Spalte X.. für Verfahrenstabelle)
ws_v.sheet_state = "hidden"


def sel(i_expr, j, model="k_model"):
    return f"INDEX({MATRIX},{i_expr},1+5*({model}-1)+{j})"


# ---------------- Ergebnisse (Calc, Fortsetzung)
r = CALC_NEXT
put(C, f"A{r}", "Ergebnis LT1 / LT2", f_bold); r += 1
res = [
    ("k_a1x", "LT1 Algorithmus x", f"={sel('k_m1', 0)}"),
    ("k_a1err", "LT1 Fehler", f"={sel('k_m1', 2)}"),
    ("k_a1warn", "LT1 Warnung", f"={sel('k_m1', 4)}"),
    ("k_a2x", "LT2 Algorithmus x", f"={sel('k_m2', 0)}"),
    ("k_a2err", "LT2 Fehler", f"={sel('k_m2', 2)}"),
    ("k_a2warn", "LT2 Warnung", f"={sel('k_m2', 4)}"),
    ("k_man1", "LT1 manuell", '=ISNUMBER(Eingabe!$I$15)'),
    ("k_man2", "LT2 manuell", '=ISNUMBER(Eingabe!$I$16)'),
    ("k_ok", "Auswertung möglich", "=AND(k_n>=4,NOT(k_dup),k_fitok)"),
    ("k_err", "Fehlermeldung", '=IF(k_n<4,"Für eine Auswertung sind mindestens 4 gültige Stufen mit Belastung und Laktat nötig (aktuell "&k_n&").",'
                               'IF(k_dup,"Zwei Stufen haben dieselbe Belastung. Bitte korrigieren.",IF(NOT(k_fitok),"Das gewählte Kurvenmodell lässt sich mit diesen Daten nicht anpassen.","")))'),
    ("k_x1", "LT1 x", '=IF(NOT(k_ok),"",IF(k_man1,MIN(k_hi,MAX(k_lo,Eingabe!$I$15)),k_a1x))'),
    ("k_x2", "LT2 x", '=IF(NOT(k_ok),"",IF(k_man2,MIN(k_hi,MAX(k_lo,Eingabe!$I$16)),k_a2x))'),
    ("k_la1", "LT1 Laktat", f'=IF(ISNUMBER(k_x1),{fsel("k_x1")},"")'),
    ("k_la2", "LT2 Laktat", f'=IF(ISNUMBER(k_x2),{fsel("k_x2")},"")'),
    ("k_hr1", "LT1 HF", f'=IF(ISNUMBER(k_x1),{hrat("k_x1")},"")'),
    ("k_hr2", "LT2 HF", f'=IF(ISNUMBER(k_x2),{hrat("k_x2")},"")'),
    ("k_zok", "Zonen möglich", "=AND(ISNUMBER(k_x1),ISNUMBER(k_x2),N(k_x2)>N(k_x1))"),
]
for nm, lab, f in res:
    put(C, f"A{r}", lab)
    C[f"B{r}"] = f
    defname(nm, "Calc", f"B{r}")
    r += 1
# Zonen
r += 1
put(C, f"A{r}", "Trainingsbereiche: Regelwert, manuell, verwendet", f_bold); r += 1
ZR0 = r
rules = ["Eingabe!$I$19/100*k_x1", "k_x1", "k_x1+Eingabe!$I$20/100*(k_x2-k_x1)", "Eingabe!$I$21/100*k_x2",
         "MAX(Eingabe!$I$21/100*k_x2,Eingabe!$I$22/100*k_hi)"]
for i in range(5):
    rr = ZR0 + i
    put(C, f"A{rr}", ["A0|A1", "A1|A2", "A2|A3", "A3|A4", "A4 Ende"][i])
    C[f"B{rr}"] = f'=IF(k_zok,{rules[i]},"")'
    C[f"C{rr}"] = f"=ISNUMBER(Eingabe!$I${26+i})"
    C[f"D{rr}"] = f'=IF(k_zok,IF(C{rr},Eingabe!$I${26+i},B{rr}),"")'
    defname(f"k_z{i}", "Calc", f"D{rr}")
    defname(f"k_zm{i}", "Calc", f"C{rr}")
r = ZR0 + 6
# Hinweise (notesFor)
put(C, f"A{r}", "Hinweise zur Auswertung (Kandidaten und laufende Nummer)", f_bold); r += 1
NT0 = r
notes = [
    'IF(AND(k_model=2,k_n<5),"Nur "&k_n&" Messpunkte für ein Polynom 3. Grades. Der kubische Fit ist mit so wenigen Punkten instabil.","")',
    'IF(AND(k_model=3,k_n<8),"Polynom 4. Grades mit "&k_n&" Punkten für 5 Parameter: Risiko künstlicher Kurvenstruktur durch Messrauschen.","")',
    'IF(k_exact,"Das Modell hat so viele Parameter wie Messpunkte und interpoliert exakt. R² = 1 ist hier bedeutungslos.","")',
    'IF(k_mind1<-0.05,"Der Fit fällt innerhalb des Messbereichs ab. Das ist physiologisch nicht plausibel; Schwellen aus Ableitungen sind dann unsicher.","")',
    'IF(AND(k_model=3,k_nwp>1),"Das Polynom 4. Grades hat "&k_nwp&" Wendepunkte im Messbereich.","")',
    'IF(k_lamax<4,"Maximales Laktat "&FIXED(k_lamax,2,TRUE)&" mmol/L: 4 mmol/L wurden nicht erreicht. Ausbelastung prüfen.","")',
    'IF(NOT(ISNUMBER(k_a1x)),"LT1 ("&Eingabe!$I$5&") nicht bestimmbar: "&k_a1err&". Anderes Verfahren wählen oder LT1 manuell setzen.","")',
    'IF(NOT(ISNUMBER(k_a2x)),"LT2 ("&Eingabe!$I$6&") nicht bestimmbar: "&k_a2err&". Anderes Verfahren wählen oder LT2 manuell setzen.","")',
    'IF(AND(ISNUMBER(k_x1),ISNUMBER(k_x2),N(k_x2)<=N(k_x1)),"LT2 liegt nicht oberhalb von LT1. Trainingsbereiche können nicht berechnet werden.","")',
    'IF(k_a1warn<>"","LT1: "&k_a1warn&".","")',
    'IF(k_a2warn<>"","LT2: "&k_a2warn&".","")',
    'IF(AND(k_zok,OR(k_z1<=k_z0,k_z2<=k_z1,k_z3<=k_z2,k_z4<=k_z3)),"Die Bereichsgrenzen sind nicht mehr aufsteigend (z. B. nach Verschieben von LT1). Bitte Grenzen anpassen oder zurücksetzen.","")',
    'IF(OR(k_man1,k_man2),"Mindestens eine Schwelle wurde manuell gesetzt. Im Bericht wird das gekennzeichnet.","")',
    'IF(AND(ISNUMBER(k_dur),N(k_dur)<3),"Stufen unter 3 min: Blutlaktat hinkt der Belastung nach, die Kurve verschiebt sich nach rechts.","")',
    'IF(NOT(k_hrok),"Zu wenige Herzfrequenzwerte für ein HF-Modell (mind. 3).","")',
]
for i, n in enumerate(notes):
    rr = NT0 + i
    C[f"B{rr}"] = f'=IF(k_ok,IFERROR({n},""),"")'
    C[f"C{rr}"] = f'=IF(B{rr}<>"",1,0)+{"C" + str(rr-1) if i else "0"}'
NT1 = NT0 + len(notes) - 1
defname("k_nnotes", "Calc", f"C{NT1}")
r = NT1 + 2
# Laktatstufen-Tabelle (gewähltes Modell)
put(C, f"A{r}", "Tabellarische Aufstellung: Laktat 2,0 … max in 0,5-Schritten (gewähltes Modell)", f_bold); r += 1
for j, h in enumerate(["L", "Anzahl", "Index", "x", "gültig", "lfd."]):
    put(C, f"{CL(1+j)}{r}", h, f_bold)
r += 1
LT0 = r
FQa, FQb, FQ = grng("Q", G0, G1 - 1), grng("Q", G0 + 1, G1), grng("Q")
for i in range(40):
    rr = LT0 + i
    C[f"A{rr}"] = 2 + 0.5 * i
    C[f"B{rr}"] = f"=SUMPRODUCT(({FQa}<A{rr})*({FQb}>=A{rr}))"
    arr(C, f"C{rr}", f"=IF(B{rr}>0,MATCH(1,({FQa}<A{rr})*({FQb}>=A{rr}),0),0)")
    fa, fb = f"INDEX({FQ},C{rr})", f"INDEX({FQ},C{rr}+1)"
    xa, xb = f"INDEX({GX},C{rr})", f"INDEX({GX},C{rr}+1)"
    C[f"D{rr}"] = f'=IF(C{rr}>0,{xa}+(A{rr}-{fa})/({fb}-{fa})*({xb}-{xa}),"")'
    C[f"E{rr}"] = f"=AND(k_ok,A{rr}<=FLOOR(k_lamax*2,1)/2+1E-9,A{rr}>=Grid!$Q${G0},A{rr}<=Grid!$Q${G1},ISNUMBER(D{rr}))"
    C[f"F{rr}"] = f'=IF(E{rr},1,0)+{"F" + str(rr-1) if i else "0"}'
LT1r = LT0 + 39
r = LT1r + 2
# Methodik-Texte (protoItems)
put(C, f"A{r}", "Protokoll und Methodik (Texte)", f_bold); r += 1
PR0 = r
xdfmt = "FIXED({},k_xd,TRUE)"
proto = [
    ("Protokoll", '=IF(k_run,"Laufband","Radergometer")&", Stufendauer "&IF(ISNUMBER(k_dur),FIXED(k_dur,1,TRUE),"–")&" min, Stufenhöhe "&IF(Eingabe!$E$5="","–",Eingabe!$E$5)&", "&k_n&" auswertbare Stufen von "&FIXED(k_lo,k_xd,TRUE)&" bis "&FIXED(k_hi,k_xd,TRUE)&" "&k_unit&"."'),
    ("Analytik", '=Eingabe!$E$6&", "&Eingabe!$E$7&IF(Eingabe!$E$8<>"",", "&Eingabe!$E$8,"")&". Für Verlaufsvergleiche immer dasselbe Messsystem und Protokoll verwenden."'),
    ("Kurvenmodell", f'=INDEX(Listen!$F$2:$F$5,k_model)&", R² "&FIXED(N(k_r2),4,TRUE)&", RMSE "&FIXED(k_rmse,3,TRUE)&" mmol/L. Das Modell ist eine mathematische Interpolation, kein physiologisches Modell des Laktatstoffwechsels."'),
    ("LT1", '=Eingabe!$I$5&IF(k_man1,", manuell auf "&FIXED(k_x1,k_xd,TRUE)&" "&k_unit&" gesetzt (Algorithmus: "&IF(ISNUMBER(k_a1x),FIXED(k_a1x,k_xd,TRUE),"–")&")","")&". Baseline für Δ-Verfahren: "&IF(Eingabe!$I$8="Ruhelaktat","Ruhelaktat","niedrigster Stufenwert")&" ("&FIXED(k_base,2,TRUE)&" mmol/L)."'),
    ("LT2", '=Eingabe!$I$6&IF(k_man2,", manuell auf "&FIXED(k_x2,k_xd,TRUE)&" "&k_unit&" gesetzt (Algorithmus: "&IF(ISNUMBER(k_a2x),FIXED(k_a2x,k_xd,TRUE),"–")&")","")&". LT2 ist eine Schätzung des oberen Übergangsbereichs und nicht per Definition MLSS oder Critical Speed/Power."'),
    ("LTP-Konzept", '=IF(OR(k_m1=1,k_m2=7),"LTP1 und LTP2 nach Hofmann/Pokan (Graz): computergestützte lineare Regressions-Breakpoint-Analyse der gemessenen Laktatwerte. LTP1 als erster Anstieg über das Ruheniveau im Suchbereich erste Stufe bis "&ROUND(k_roi*100,0)&" % vmax, LTP2 als zweiter, beschleunigter Anstieg im Bereich LTP1 bis vmax. Daraus ergibt sich das Drei-Phasen-Modell: Phase I unterhalb LTP1, Phase II zwischen LTP1 und LTP2, Phase III oberhalb LTP2.","")'),
    ("Hilfsgeraden", '=IF(k_gfree,"Drei stetig verbundene Regressionsgeraden (Grundlinie, Übergang, steiler Anstieg) mit Knoten an LT1 und LT2, nach der Methode der kleinsten Quadrate an die Messpunkte angepasst. Das LTP-Verfahren wählt die Knoten mit dem kleinsten Fehler.","Drei verbundene Geraden (Grundlinie, Übergang, steiler Anstieg) mit Eckpunkten auf der Laktatkurve bei Testbeginn, LT1, LT2 und Testende. Wo die Kurve nicht konvex ist, wird das betroffene Geradenstück angehoben, sodass die Hilfsgeraden die Kurve nirgends unterschreiten. Die Knoten der LTP-Verfahren stammen weiterhin aus der freien 3-Phasen-Regression durch die Messpunkte; die Korrektur betrifft nur die Darstellung der Geraden.")'),
    ("Dmax-Verfahren", '="Gesucht wird der Punkt, an dem die Tangente der Kurve parallel zur Sehne liegt (f′(x) = m), d. h. der größte Abstand der Kurve unterhalb der Sehne. In Excel wird dazu die Kurve in 1000 Schritten ausgewertet und das Maximum per Parabel verfeinert. Das Ergebnis ist unabhängig von der Achsenskalierung."'),
    ("Trainingsbereiche", '="A1 ab "&Eingabe!$I$19&" % LT1, A1|A2 "&IF(k_zm1,"manuell gesetzt","bei LT1")&", A2 bis "&Eingabe!$I$20&" % der Strecke LT1→LT2, A3 bis "&Eingabe!$I$21&" % LT2, A4 bis "&Eingabe!$I$22&" % vmax. HF-Werte aus dem "&IF(k_hrdeg=1,"linearen","quadratischen")&" HF-Modell."&IF(OR(k_zm0,k_zm1,k_zm2,k_zm3,k_zm4)," Manuell angepasste Grenzen: "&MID(IF(k_zm0,", A0|A1","")&IF(k_zm1,", A1|A2","")&IF(k_zm2,", A2|A3","")&IF(k_zm3,", A3|A4","")&IF(k_zm4,", A4 Ende",""),3,200)&".","")'),
]
for i, (k, f) in enumerate(proto):
    rr = PR0 + i
    put(C, f"A{rr}", k)
    C[f"B{rr}"] = f'=IF(k_ok,IFERROR({f[1:]},""),"")'
PR1 = PR0 + len(proto) - 1
r = PR1 + 2
# Alter
put(C, f"A{r}", "Alter [Jahre]")
C[f"B{r}"] = '=IF(AND(ISNUMBER(Eingabe!$B$5),ISNUMBER(Eingabe!$B$6)),(Eingabe!$B$6-Eingabe!$B$5)/365.2425,IF(ISNUMBER(Eingabe!$B$5),(TODAY()-Eingabe!$B$5)/365.2425,""))'
defname("k_age", "Calc", f"B{r}")
r += 1
C.column_dimensions["A"].width = 34
ws_c.sheet_state = "hidden"

# ====================================================================== Guide (Hilfsgeraden)
D = ws_gd
put(D, "A1", "Hilfsgeraden (3 Phasen): Knoten an LT1/LT2 (sonst Regressions-Optimum); Modus korrigiert = Sehnen, angehoben wo nötig", f_bold)
D["A2"] = "Knoten an LT1/LT2 verwendbar"
D["B2"] = "=AND(k_ok,ISNUMBER(k_x1),ISNUMBER(k_x2),N(k_x2)>N(k_x1),N(k_x1)>k_lo,N(k_x2)<k_hi)"
D["A3"] = "Hilfsgeraden vorhanden"; D["B3"] = "=AND(k_ok,OR(B2,k_3ok))"
D["A4"] = "Knoten 1"; D["B4"] = "=IF(B2,k_x1,IF(k_3ok,k_3k1,k_lo+k_rng/3))"
D["A5"] = "Knoten 2"; D["B5"] = "=IF(B2,k_x2,IF(k_3ok,k_3k2,k_lo+2*k_rng/3))"
D["A6"] = "an Regressions-Optimum"; D["B6"] = "=AND(k_3ok,ABS(B4-k_3k1)<1E-6,ABS(B5-k_3k2)<1E-6)"
D["A8"] = "Y0 (lo)"; D["A9"] = "Y1 (Knoten 1)"; D["A10"] = "Y2 (Knoten 2)"; D["A11"] = "Y3 (hi)"
D["A12"] = "angehoben um"
# freie Hinge-Regression an den Knoten
put(D, "D2", "freie Regression an den Knoten", f_bold)
evg, _ = hinge_eval(D, "D", "E", 3, "$B$4", "$B$5", "Guide")
put(D, "G2", "Sehnen-Konstruktion", f_bold)
# Stützstellen je Segment: 61 Punkte
put(D, "G3", "j"); put(D, "H3", "f Seg 1"); put(D, "I3", "f Seg 2"); put(D, "J3", "f Seg 3")
Xk = ["k_lo", "$B$4", "$B$5", "k_hi"]
for j in range(61):
    rr = 4 + j
    D[f"G{rr}"] = j
    for s in range(3):
        xx = f"({Xk[s]}+({Xk[s+1]}-{Xk[s]})*G{rr}/60)"
        D[f"{'HIJ'[s]}{rr}"] = f"={fsel(xx)}"
SJ = "$G$4:$G$64"
# 36 Schritte (12 Iterationen × 3 Segmente)
put(D, "L3", "Schritt"); put(D, "M3", "Seg"); put(D, "N3", "d"); put(D, "O3", "Y0"); put(D, "P3", "Y1"); put(D, "Q3", "Y2"); put(D, "R3", "Y3")
D["O4"] = f"={fsel('k_lo')}"
D["P4"] = f"={fsel('$B$4')}"
D["Q4"] = f"={fsel('$B$5')}"
D["R4"] = f"={fsel('k_hi')}"
D["L4"] = 0
for s in range(36):
    rr = 5 + s
    seg = s % 3
    D[f"L{rr}"] = s + 1
    D[f"M{rr}"] = seg + 1
    Y = ["O", "P", "Q", "R"]
    yi, yj = f"{Y[seg]}{rr-1}", f"{Y[seg+1]}{rr-1}"
    fcol = f"${'HIJ'[seg]}$4:${'HIJ'[seg]}$64"
    arr(D, f"N{rr}", f"=IF({Xk[seg+1]}>{Xk[seg]},MAX(0,MAX({fcol}-({yi}+({yj}-{yi})*{SJ}/60))),0)")
    for q in range(4):
        prev = f"{Y[q]}{rr-1}"
        if q in (seg, seg + 1):
            D[f"{Y[q]}{rr}"] = f"=IF(N{rr}>1E-7,{prev}+N{rr},{prev})"
        else:
            D[f"{Y[q]}{rr}"] = f"={prev}"
LASTG = 5 + 35
for i, col in enumerate("OPQR"):
    D[f"B{8+i}"] = f"=IF(k_gfree,{['fL', 'f1', 'f2', 'fH'][i]},{col}{LASTG})"
# freie Variante: Werte an lo, k1, k2, hi
fL = f"({evg['c0']}+{evg['bx']}*k_lo)"
f1 = f"({evg['c0']}+{evg['bx']}*$B$4)"
f2 = f"({evg['c0']}+{evg['bx']}*$B$5+{evg['b1']}*($B$5-$B$4))"
fH = f"({evg['c0']}+{evg['bx']}*k_hi+{evg['b1']}*(k_hi-$B$4)+{evg['b2']}*(k_hi-$B$5))"
for i, v in enumerate([fL, f1, f2, fH]):
    D[f"B{8+i}"] = f"=IF(k_gfree,{v},{'OPQR'[i]}{LASTG})"
D["B12"] = f"=MAX(0,O{LASTG}-O4,P{LASTG}-P4,Q{LASTG}-Q4,R{LASTG}-R4)"
ws_gd.sheet_state = "hidden"

# ====================================================================== ChartData
CD = ws_cd
put(CD, "A1", "Diagrammdaten (automatisch)", f_bold)
# Laktatkurve 201 Punkte
CD["A2"] = "x"; CD["B2"] = "La Modell"; CD["C2"] = "HF Modell"; CD["D2"] = "La/x Modell"; CD["E2"] = "x HF"
CD["F2"] = "ln-Achse x"; CD["G2"] = "x Glu"; CD["H2"] = "Glu Interp."
for i in range(201):
    rr = 3 + i
    gi = G0 + 5 * i
    CD[f"A{rr}"] = f"=IF(k_ok,Grid!$B${gi},NA())"
    CD[f"B{rr}"] = f"=IF(k_ok,Grid!$Q${gi},NA())"
    CD[f"E{rr}"] = f"=IF(k_hrok,k_hx0+(k_hx1-k_hx0)*{i}/200,NA())"
    CD[f"C{rr}"] = f"=IF(k_hrok,k_h0+k_h1*E{rr}+k_h2*E{rr}^2,NA())"
    CD[f"D{rr}"] = f"=IF(k_ok,Grid!$Q${gi}/Grid!$B${gi},NA())"
    # Glukose monotone kubische Hermite-Interpolation
    x0g, x1g = f"Calc!$AK${C0}", f"INDEX(Calc!$AK${C0}:$AK${C1},k_ng)"
    CD[f"G{rr}"] = f"=IF(AND(k_gluOn,k_ng>=2),{x0g}+({x1g}-{x0g})*{i}/200,NA())"
    idx = f"MIN(k_ng-1,MAX(1,MATCH(G{rr},Calc!$AK${C0}:INDEX(Calc!$AK${C0}:$AK${C1},k_ng),1)))"
    CD[f"I{rr}"] = f"=IF(AND(k_gluOn,k_ng>=2),{idx},1)"
    ii = f"I{rr}"
    xa = f"INDEX(Calc!$AK${C0}:$AK${C1},{ii})"
    xb = f"INDEX(Calc!$AK${C0}:$AK${C1},{ii}+1)"
    ya = f"INDEX(Calc!$AL${C0}:$AL${C1},{ii})"
    yb = f"INDEX(Calc!$AL${C0}:$AL${C1},{ii}+1)"
    ma = f"INDEX(Calc!$AP${C0}:$AP${C1},{ii})"
    mb = f"INDEX(Calc!$AQ${C0}:$AQ${C1},{ii})"
    h = f"({xb}-{xa})"
    t = f"((G{rr}-{xa})/{h})"
    CD[f"H{rr}"] = (f"=IF(AND(k_gluOn,k_ng>=2),(2*{t}^3-3*{t}^2+1)*{ya}+({t}^3-2*{t}^2+{t})*{h}*{ma}"
                    f"+(-2*{t}^3+3*{t}^2)*{yb}+({t}^3-{t}^2)*{h}*{mb},NA())")
# vertikale Linien und Hilfsserien ab Spalte K
CD["K2"] = "Serie"; CD["L2"] = "x"; CD["M2"] = "y"
lines = {}


def vline(key, xexpr, y0, y1, label, row):
    CD[f"K{row}"] = label
    CD[f"L{row}"] = f'=IF(ISNUMBER({xexpr}),{xexpr},NA())'
    CD[f"M{row}"] = f"=IF(ISNUMBER(L{row}),{y0},NA())"
    CD[f"L{row+1}"] = f'=IF(ISNUMBER({xexpr}),{xexpr},NA())'
    CD[f"M{row+1}"] = f"=IF(ISNUMBER(L{row+1}),{y1},NA())"
    lines[key] = (row, row + 1)


CD["O2"] = "yMax La"; CD["O3"] = "=CEILING(MAX(k_lamax*1.1,4.5),1)"
CD["P2"] = "HF min"; CD["P3"] = '=IF(k_nh>0,FLOOR(MIN(OFFSET(Calc!$AF$3,0,0,MAX(k_nh,1),1))-10,10),0)'
CD["Q2"] = "HF max"; CD["Q3"] = '=IF(k_nh>0,CEILING(k_hrmax+6,10),200)'
CD["R2"] = "LE min"; CD["S2"] = "LE max"
arr(CD, "R3", f"=MIN(IF({RM}=1,{RY}/{RX}))*0.9")
arr(CD, "S3", f"=MAX(IF({RM}=1,{RY}/{RX}))*1.1")
row = 3
for key, xe, y0, y1, lab in (("lt1", "k_x1", 0, "$O$3", "LT1"), ("lt2", "k_x2", 0, "$O$3", "LT2"),
                             ("z0", "k_z0", 0, "$O$3", "A0|A1"), ("z2", "k_z2", 0, "$O$3", "A2|A3"),
                             ("z3", "k_z3", 0, "$O$3", "A3|A4"), ("z4", "k_z4", 0, "$O$3", "A4 Ende"),
                             ("z1", 'IF(k_zm1,k_z1,"")', 0, "$O$3", "A1|A2"),
                             ("hlt1", "k_x1", "$P$3", "$Q$3", "LT1 HF"), ("hlt2", "k_x2", "$P$3", "$Q$3", "LT2 HF"),
                             ("elt1", "k_x1", "$R$3", "$S$3", "LT1 LE"), ("elt2", "k_x2", "$R$3", "$S$3", "LT2 LE")):
    vline(key, xe, y0, y1, lab, row)
    row += 3
# Hilfsgeraden (4 Punkte)
CD[f"K{row}"] = "Hilfsgeraden"
GR = row
for i, (xe, ye) in enumerate((("k_gx0", "k_gy0"), ("k_gk1", "k_gy1"), ("k_gk2", "k_gy2"), ("k_gx3", "k_gy3"))):
    CD[f"L{row+i}"] = f"=IF(k_gok,{xe},NA())"
    CD[f"M{row+i}"] = f"=IF(k_gok,{ye},NA())"
row += 5
# Dmax-Sehne des gewählten LT2-Verfahrens
CD[f"K{row}"] = "Dmax-Sehne"
CH = row
chord_src = {MID["dmax"]: DM, MID["moddmax"]: MD, MID["lpdmax"]: LPD, MID["expdmax"]: EXD}
cond = "OR(" + ",".join(f"k_m2={i+1}" for i in chord_src) + ")"


def pick(field):
    expr = "0"
    for i, b in chord_src.items():
        expr = f"IF(k_m2={i+1},INDEX(VCalc!$C${b[field]}:$F${b[field]},k_model),{expr})"
    return expr


CD[f"L{row}"] = f'=IF(AND(k_ok,{cond},ISNUMBER(k_a2x)),{pick("xs")},NA())'
CD[f"M{row}"] = f'=IF(AND(k_ok,{cond},ISNUMBER(k_a2x)),{pick("ys")},NA())'
CD[f"L{row+1}"] = f'=IF(AND(k_ok,{cond},ISNUMBER(k_a2x)),{pick("xe")},NA())'
CD[f"M{row+1}"] = f'=IF(AND(k_ok,{cond},ISNUMBER(k_a2x)),{pick("ye")},NA())'
row += 3
# Log-Log Regressionsgeraden
CD[f"K{row}"] = "Log-Log Geraden"
LLR = row
pts = [("k_lo", "k_llal", "k_llbl"), ("k_llx", "k_llal", "k_llbl"), ("k_llx", "k_llar", "k_llbr"), (lastx, "k_llar", "k_llbr")]
for i, (xe, a, b) in enumerate(pts):
    CD[f"L{row+i}"] = f"=IF(k_llok,{xe},NA())"
    CD[f"M{row+i}"] = f"=IF(k_llok,EXP({a}+{b}*LN({xe})),NA())"
row += 5
# LEmin-Punkt
CD[f"K{row}"] = "LEmin"
LEP = row
CD[f"L{row}"] = f'=IF(k_ok,INDEX(VCalc!$C${LEM["x"]}:$F${LEM["x"]},k_model),NA())'
CD[f"M{row}"] = f'=IF(k_ok,INDEX(VCalc!$C${LEM["la"]}:$F${LEM["la"]},k_model)/L{row},NA())'
row += 2
# Ruheglukose
CD[f"K{row}"] = "Ruheglukose"
RG = row
CD[f"L{row}"] = f'=IF(AND(k_gluOn,k_ng>=2,ISNUMBER(Eingabe!$E$10)),Calc!$AK${C0},NA())'
CD[f"M{row}"] = f'=IF(ISNUMBER(L{row}),Eingabe!$E$10,NA())'
CD[f"L{row+1}"] = f'=IF(AND(k_gluOn,k_ng>=2,ISNUMBER(Eingabe!$E$10)),INDEX(Calc!$AK${C0}:$AK${C1},k_ng),NA())'
CD[f"M{row+1}"] = f'=IF(ISNUMBER(L{row+1}),Eingabe!$E$10,NA())'
row += 3
ws_cd.sheet_state = "hidden"

# ====================================================================== Diagramme (Objekte)


def mk_series(ws, xc, yc, r0, r1, title, color, marker=False, dash=None, width=2.0, smooth=False):
    xs = Reference(ws, min_col=xc, min_row=r0, max_row=r1)
    ys = Reference(ws, min_col=yc, min_row=r0, max_row=r1)
    s = Series(ys, xs, title=title)
    s.smooth = smooth
    if marker:
        s.marker.symbol = "circle"
        s.marker.size = 7
        s.marker.graphicalProperties.solidFill = color
        s.marker.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.noFill = True
    else:
        s.marker.symbol = "none"
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = int(width * 12700)
        if dash:
            s.graphicalProperties.line.dashStyle = dash
    return s


def base_chart(title, xt, yt, w=17, h=9.5):
    ch = ScatterChart(scatterStyle="lineMarker")
    ch.title = title or None
    ch.style = 2
    ch.x_axis.title = xt
    ch.y_axis.title = yt
    ch.width, ch.height = w, h
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"
    ch.x_axis.majorGridlines = None
    return ch


XC, YC = 12, 13   # ChartData Spalten L/M


def la_chart(title="Laktat zu Belastung"):
    ch = base_chart(title, "Belastung", "Laktat [mmol/L]")
    ch.series.append(mk_series(ws_cd, 1, 2, 3, 203, "Laktatkurve (Modell)", INK, width=2.25))
    ch.series.append(mk_series(ws_c, 22, 23, C0, C1, "Messwerte", INK, marker=True))
    ch.series.append(mk_series(ws_cd, XC, YC, GR, GR + 3, "Hilfsgeraden", "E5482A", width=1.5))
    ch.series.append(mk_series(ws_cd, XC, YC, CH, CH + 1, "Dmax-Sehne", MUTED, dash="dash", width=1))
    r0, r1 = lines["lt1"]
    ch.series.append(mk_series(ws_cd, XC, YC, r0, r1, "LT1", LT1C, width=2))
    r0, r1 = lines["lt2"]
    ch.series.append(mk_series(ws_cd, XC, YC, r0, r1, "LT2", LT2C, width=2))
    for k, colr in (("z0", "7FB58A"), ("z1", "5E9E73"), ("z2", "6FA29A"), ("z3", "C9A227"), ("z4", "D9725E")):
        r0, r1 = lines[k]
        ch.series.append(mk_series(ws_cd, XC, YC, r0, r1, CD[f"K{r0}"].value, colr, dash="dash", width=1))
    ch.y_axis.scaling.min = 0
    return ch


def hr_chart(title="Herzfrequenz zu Belastung"):
    ch = base_chart(title, "Belastung", "HF [1/min]")
    ch.y_axis.scaling.min = 60
    ch.series.append(mk_series(ws_cd, 5, 3, 3, 203, "HF-Modell", INK, width=2.25))
    ch.series.append(mk_series(ws_c, 33, 34, C0, C1, "HF gemessen", INK, marker=True))
    for k, colr, lab in (("hlt1", LT1C, "LT1"), ("hlt2", LT2C, "LT2")):
        r0, r1 = lines[k]
        ch.series.append(mk_series(ws_cd, XC, YC, r0, r1, lab, colr, width=2))
    return ch


# ====================================================================== Auswertung
O = ws_out
for col, w in zip("ABCDEFGHI", (34, 23, 23, 15, 15, 11, 11, 11, 12)):
    O.column_dimensions[col].width = w
O["A1"] = '=IF(Eingabe!$B$4<>"",Eingabe!$B$4&IF(ISNUMBER(Eingabe!$B$6)," · Test vom "&TEXT(Eingabe!$B$6,"DD.MM.YYYY"),""),"Laktatdiagnostik – Auswertung")'
O["A1"].font = f_head
put(O, "A2", "Stufentest-Auswertung mit LT1/LT2, Modellvergleich und Trainingsbereichen. Alle Werte berechnen sich aus dem Blatt „Eingabe“.", f_muted)
O["A3"] = '=IF(k_ok,"",k_err)'
O["A3"].font = Font(name=FONT, size=11, bold=True, color="A1161F")
O.merge_cells("A3:I3")
# Schwellenkarten
put(O, "A5", "Schwellen", f_h2)
hdr = ["", "LT1", "LT2", "Maximum"]
for j, h in enumerate(hdr):
    c = O.cell(6, 1 + j, h)
    c.font = Font(name=FONT, size=11, bold=True, color=[INK, LT1C, LT2C, INK][j])
    c.border = Border(bottom=Side(style="medium", color=INK))
    c.alignment = Alignment(horizontal="left" if j == 0 else "right")
rows_thr = [
    ("Verfahren", '=IF(k_ok,Eingabe!$I$5&IF(k_man1," (manuell)",""),"–")', '=IF(k_ok,Eingabe!$I$6&IF(k_man2," (manuell)",""),"–")', '="letzte Stufe"'),
    ('="Belastung ["&k_unit&"]"', '=IF(ISNUMBER(k_x1),ROUND(k_x1,k_xd),"–")', '=IF(ISNUMBER(k_x2),ROUND(k_x2,k_xd),"–")', '=IF(k_ok,ROUND(k_hi,k_xd),"–")'),
    ("Pace [min/km]", f'=IF(k_run,{pace("k_x1")},"–")', f'=IF(k_run,{pace("k_x2")},"–")', f'=IF(AND(k_ok,k_run),{pace("k_hi")},"–")'),
    ("Geschwindigkeit [m/s]", '=IF(AND(k_run,ISNUMBER(k_x1)),ROUND(k_x1/3.6,2),"–")', '=IF(AND(k_run,ISNUMBER(k_x2)),ROUND(k_x2/3.6,2),"–")', '=IF(AND(k_ok,k_run),ROUND(k_hi/3.6,2),"–")'),
    ("Laktat [mmol/L]", '=IF(ISNUMBER(k_la1),ROUND(k_la1,2),"–")', '=IF(ISNUMBER(k_la2),ROUND(k_la2,2),"–")', '=IF(k_ok,k_lamax,"–")'),
    ("Herzfrequenz [1/min]", '=IF(ISNUMBER(k_hr1),ROUND(k_hr1,0),"–")', '=IF(ISNUMBER(k_hr2),ROUND(k_hr2,0),"–")', '=IF(ISNUMBER(k_hrmax),k_hrmax,"–")'),
    ("% HFmax", '=IF(AND(ISNUMBER(k_hr1),ISNUMBER(k_hrmax)),ROUND(k_hr1/k_hrmax*100,1),"–")', '=IF(AND(ISNUMBER(k_hr2),ISNUMBER(k_hrmax)),ROUND(k_hr2/k_hrmax*100,1),"–")', '=IF(ISNUMBER(k_hrmax),100,"–")'),
    ("Anteil vmax [%]", '=IF(ISNUMBER(k_x1),ROUND(k_x1/k_hi*100,1),"–")', '=IF(ISNUMBER(k_x2),ROUND(k_x2/k_hi*100,1),"–")', '=IF(k_ok,100,"–")'),
    ("Ruhelaktat [mmol/L]", '="–"', '="–"', '=IF(ISNUMBER(k_rest),k_rest,"–")'),
    ('="LT2 − LT1 ["&k_unit&"]"', '="–"', '="–"', '=IF(AND(ISNUMBER(k_x1),ISNUMBER(k_x2)),ROUND(k_x2-k_x1,k_xd),"–")'),
    ("Algorithmus (vor manueller Korrektur)", '=IF(ISNUMBER(k_a1x),ROUND(k_a1x,k_xd),"–")', '=IF(ISNUMBER(k_a2x),ROUND(k_a2x,k_xd),"–")', '=""'),
]
for i, rowv in enumerate(rows_thr):
    rr = 7 + i
    for j, v in enumerate(rowv):
        c = O.cell(rr, 1 + j, v)
        c.font = f_bold if (j and i == 1) else f_base
        if j and i == 1:
            c.font = Font(name=FONT, size=14, bold=True, color=[INK, LT1C, LT2C, INK][j])
        c.alignment = Alignment(horizontal="left" if j == 0 else "right")
        c.border = b_bottom
THR_END = 7 + len(rows_thr)
put(O, f"A{THR_END}", "Manuelle Korrektur von LT1/LT2 und Bereichsgrenzen im Blatt „Eingabe“ (Spalte I).", f_muted)
# Hinweise
NH = THR_END + 2
put(O, f"A{NH}", "Hinweise zur Auswertung", f_h2)
for i in range(len(notes)):
    rr = NH + 1 + i
    O[f"A{rr}"] = f'=IFERROR(INDEX(Calc!$B${NT0}:$B${NT1},MATCH({i+1},Calc!$C${NT0}:$C${NT1},0)),"")'
    O[f"A{rr}"].font = Font(name=FONT, size=10, color="9A5B00")
O.conditional_formatting.add(f"A{NH+1}:A{NH+len(notes)}",
                             FormulaRule(formula=[f'OR(LEFT(A{NH+1},4)="LT1 ",LEFT(A{NH+1},4)="LT2 ",LEFT(A{NH+1},5)="Die B")'],
                                         font=Font(name=FONT, color="A1161F", bold=True)))
# Trainingsbereiche
ZT = NH + len(notes) + 2
put(O, f"A{ZT}", "Trainingsbereiche", f_h2)
zh = ["Trainingsbereich", "HF von", "HF bis", '="Belastung von ["&k_unit&"]"', '="Belastung bis ["&k_unit&"]"',
      "Pace von", "Pace bis", "Laktat bis", "manuell"]
for j, h in enumerate(zh):
    c = O.cell(ZT + 1, 1 + j, h)
    c.font = f_bold
    c.border = Border(bottom=Side(style="medium", color=INK))
    c.alignment = Alignment(horizontal="left" if j == 0 else "right", wrap_text=True)
ZNAMES = ["A0 · Regeneration/Kompensation", "A1 · Extensives Ausdauertraining 1", "A2 · Extensives Ausdauertraining 2",
          "A3 · Intensives Ausdauertraining", "A4 · Entwicklungsbereich"]
for i in range(5):
    rr = ZT + 2 + i
    fr = f"k_z{i-1}" if i else '""'
    to = f"k_z{i}"
    vals = [ZNAMES[i],
            f'=IF(AND(k_zok,{i}>0),IFERROR(ROUND({hrat(fr)},0),"–"),"")' if i else '=""',
            f'=IF(k_zok,IFERROR(ROUND({hrat(to)},0),"–"),"–")',
            f'=IF(AND(k_zok,{i}>0),ROUND({fr},k_xd),"")' if i else '=""',
            f'=IF(k_zok,ROUND({to},k_xd),"–")',
            (f'=IF(AND(k_zok,k_run),{pace(fr)},"")' if i else '=""'),
            f'=IF(AND(k_zok,k_run),{pace(to)},"")',
            f'=IF(k_zok,ROUND({fsel(f"MIN({to},k_hi)")},2),"–")',
            f'=IF(k_zm{i},"• manuell",IF({i}=1,"folgt LT1",""))']
    for j, v in enumerate(vals):
        c = O.cell(rr, 1 + j, v)
        c.font = f_bold if j == 0 else f_base
        c.fill = PatternFill("solid", fgColor=ZC[i])
        c.alignment = Alignment(horizontal="left" if j == 0 else "right")
        if j in (1, 2):
            c.number_format = "0"
        c.border = b_bottom
put(O, f"A{ZT+7}", "Grenzen: A1 ab % LT1, A1|A2 = LT1, A2 bis % der Strecke LT1→LT2, A3 bis % LT2, A4 bis % vmax (Regeln im Blatt „Eingabe“). HF aus dem HF-Modell.", f_muted)
# Laktatstufen
LVT = ZT + 9
put(O, f"A{LVT}", "Tabellarische Aufstellung", f_h2)
for j, h in enumerate(["Laktat [mmol/L]", '="Belastung ["&k_unit&"]"', "Pace [min/km]", "Belastung [m/s]", "HF [1/min]"]):
    c = O.cell(LVT + 1, 1 + j, h)
    c.font = f_bold
    c.border = Border(bottom=Side(style="medium", color=INK))
    c.alignment = Alignment(horizontal="left" if j == 0 else "right")
for i in range(30):
    rr = LVT + 2 + i
    m = f"MATCH({i+1},Calc!$F${LT0}:$F${LT1r},0)"
    O[f"A{rr}"] = f'=IFERROR(INDEX(Calc!$A${LT0}:$A${LT1r},{m}),"")'
    O[f"B{rr}"] = f'=IFERROR(ROUND(INDEX(Calc!$D${LT0}:$D${LT1r},{m}),k_xd),"")'
    O[f"C{rr}"] = f'=IF(AND(k_run,ISNUMBER(B{rr})),{pace(f"INDEX(Calc!$D${LT0}:$D${LT1r},{m})")},"")'
    O[f"D{rr}"] = f'=IF(AND(k_run,ISNUMBER(B{rr})),ROUND(INDEX(Calc!$D${LT0}:$D${LT1r},{m})/3.6,2),"")'
    O[f"E{rr}"] = f'=IF(ISNUMBER(B{rr}),IFERROR(ROUND({hrat(f"INDEX(Calc!$D${LT0}:$D${LT1r},{m})")},0),"–"),"")'
    for cc in "ABCDE":
        O[f"{cc}{rr}"].font = f_base
        O[f"{cc}{rr}"].alignment = Alignment(horizontal="left" if cc == "A" else "right")
    O[f"A{rr}"].number_format = "0.0"
O.add_chart(la_chart(), "K5")
O.add_chart(hr_chart(), "K26")
O.freeze_panes = "A4"

# ====================================================================== Verfahren
W = ws_meth
for col, w in zip("ABCDEFGH", (30, 28, 12, 10, 8, 62, 10, 10)):
    W.column_dimensions[col].width = w
put(W, "A1", "Alle Schwellenverfahren", f_head)
put(W, "A2", "Berechnet mit dem gewählten Kurvenmodell. Keine Extrapolation über den Messbereich hinaus. LT1/LT2-Auswahl im Blatt „Eingabe“.", f_muted)
heads = ["Verfahren", "Zielkonstrukt", '="Belastung ["&k_unit&"]"', "Laktat", "HF", "Hinweis", "als LT1", "als LT2"]
for j, h in enumerate(heads):
    c = W.cell(4, 1 + j, h)
    c.font = f_bold
    c.border = Border(bottom=Side(style="medium", color=INK))
    c.alignment = Alignment(horizontal="left" if j in (0, 1, 5) else "right")
rr = 5
grp = None
for i, (mid, name, g, ctr) in enumerate(METHODS):
    if g != grp:
        grp = g
        c = W.cell(rr, 1, GRP_NAMES[g])
        c.font = f_bold
        c.border = Border(bottom=Side(style="thin", color=INK))
        rr += 1
    xi = sel(i + 1, 0)
    W[f"A{rr}"] = name
    W[f"B{rr}"] = ctr
    W[f"C{rr}"] = f'=IF(AND(k_ok,ISNUMBER({xi})),ROUND({xi},k_xd),"–")'
    W[f"D{rr}"] = f'=IF(AND(k_ok,ISNUMBER({xi})),ROUND({sel(i+1, 1)},2),"–")'
    W[f"E{rr}"] = f'=IF(AND(k_ok,ISNUMBER({xi})),IFERROR(ROUND({hrat(xi)},0),"–"),"–")'
    W[f"F{rr}"] = (f'=IF(k_ok,MID(IF({sel(i+1, 2)}<>""," · "&{sel(i+1, 2)},"")&IF({sel(i+1, 3)}<>""," · "&{sel(i+1, 3)},"")'
                   f'&IF({sel(i+1, 4)}<>""," · "&{sel(i+1, 4)},""),4,500),"")')
    W[f"G{rr}"] = f'=IF(k_m1={i+1},"● LT1","")'
    W[f"H{rr}"] = f'=IF(k_m2={i+1},"● LT2","")'
    for cc in "ABCDEFGH":
        W[f"{cc}{rr}"].font = f_base
        W[f"{cc}{rr}"].border = b_bottom
        W[f"{cc}{rr}"].alignment = Alignment(horizontal="left" if cc in "ABF" else "right", wrap_text=(cc == "F"),
                                            vertical="top")
    W[f"B{rr}"].font = f_muted
    W[f"F{rr}"].font = f_muted
    W[f"G{rr}"].font = Font(name=FONT, size=10, bold=True, color=LT1C)
    W[f"H{rr}"].font = Font(name=FONT, size=10, bold=True, color=LT2C)
    W.conditional_formatting.add(f"A{rr}:F{rr}", FormulaRule(formula=[f"OR(k_m1={i+1},k_m2={i+1})"], fill=fill_soft,
                                                             font=Font(name=FONT, bold=True)))
    rr += 1
MC_T = rr + 2
put(W, f"A{MC_T}", "Modellvergleich am selben Datensatz", f_h2)
put(W, f"A{MC_T+1}", "Wie stark sich LT1 und LT2 allein durch das Kurvenmodell verschieben. Ein höheres R² macht ein Modell nicht besser.", f_muted)
mh = ["Modell", "Parameter", "R²", "RMSE", '="LT1 · "&Eingabe!$I$5', '="LT2 · "&Eingabe!$I$6', "Δ LT2"]
for j, h in enumerate(mh):
    c = W.cell(MC_T + 2, 1 + j, h)
    c.font = f_bold
    c.border = Border(bottom=Side(style="medium", color=INK))
    c.alignment = Alignment(horizontal="left" if j == 0 else "right")
for m in range(4):
    rr = MC_T + 3 + m
    q = QF + 1 + m
    ok = f"Calc!$B${q}"
    l1 = f"INDEX({MATRIX},k_m1,{1+5*m})"
    l2 = f"INDEX({MATRIX},k_m2,{1+5*m})"
    ref2 = f"INDEX({MATRIX},k_m2,1+5*(k_model-1))"
    W[f"A{rr}"] = f'="{MODEL_SHORT[m]}"&IF(AND({ok},Calc!$G${q})," (exakt)","")&IF(k_model={m+1},"  ◄ gewählt","")'
    W[f"B{rr}"] = f'=IF({ok},Calc!$C${q},"nicht anpassbar")'
    W[f"C{rr}"] = f'=IF(AND(k_ok,{ok}),ROUND(Calc!$E${q},4),"–")'
    W[f"D{rr}"] = f'=IF(AND(k_ok,{ok}),ROUND(Calc!$F${q},3),"–")'
    W[f"E{rr}"] = f'=IF(AND(k_ok,{ok},ISNUMBER({l1})),ROUND({l1},k_xd),"–")'
    W[f"F{rr}"] = f'=IF(AND(k_ok,{ok},ISNUMBER({l2})),ROUND({l2},k_xd),"–")'
    W[f"G{rr}"] = (f'=IF(AND(k_ok,{ok},ISNUMBER({l2}),ISNUMBER({ref2})),IF(k_model={m+1},"Referenz",'
                   f'IF({l2}-{ref2}>=0,"+","")&FIXED({l2}-{ref2},k_xd,TRUE)),"–")')
    for j in range(7):
        c = W.cell(rr, 1 + j)
        c.font = f_base
        c.border = b_bottom
        c.alignment = Alignment(horizontal="left" if j == 0 else "right")
    W.conditional_formatting.add(f"A{rr}:G{rr}", FormulaRule(formula=[f"k_model={m+1}"], fill=fill_soft,
                                                             font=Font(name=FONT, bold=True)))
W[f"A{MC_T+8}"] = '=IF(k_ok,"Koeffizienten des gewählten Modells (x in "&k_unit&"): "&k_eq,"")'
W[f"A{MC_T+8}"].font = f_muted
W[f"A{MC_T+9}"] = '=IF(k_ok,"Güte gewählt: R² "&FIXED(N(k_r2),4,TRUE)&" · RMSE "&FIXED(k_rmse,3,TRUE)&" mmol/L · n = "&k_n&IF(k_gok," · Hilfsgeraden-RMSE "&FIXED(k_grmse,3,TRUE)&" mmol/L"&IF(Guide!$B$6," (optimale Knoten)"," (Knoten an LT1/LT2) · Regressions-Optimum bei "&IF(k_3ok,FIXED(k_3k1,k_xd,TRUE)&" / "&FIXED(k_3k2,k_xd,TRUE),"–")),""),"")'
W[f"A{MC_T+9}"].font = f_muted
W.freeze_panes = "A5"

# ====================================================================== Diagramme
DG = ws_dia
put(DG, "A1", "Weitere Darstellungen", f_head)
put(DG, "A2", "Laktatäquivalent: Das Minimum markiert den Übergang aus dem aeroben Grundbereich (Grundlage des Dickhuth-Verfahrens: LEmin + 1,5 mmol/L). "
    "Log-Log: ln Laktat gegen ln Belastung, zwei Regressionsgeraden, Schnittpunkt = Log-Log-Breakpoint.", f_muted)
ch = base_chart("Laktatäquivalent (Laktat ÷ Belastung)", "Belastung", "Laktat ÷ Belastung", 16, 9)
ch.series.append(mk_series(ws_cd, 1, 4, 3, 203, "La/x Modell", INK, width=2))
ch.series.append(mk_series(ws_c, 22, 24, C0, C1, "Messwerte", INK, marker=True))
s = mk_series(ws_cd, XC, YC, LEP, LEP, "LEmin", ACC, marker=True)
s.marker.symbol = "triangle"
s.marker.graphicalProperties.solidFill = INK
ch.series.append(s)
for k, colr, lab in (("elt1", LT1C, "LT1"), ("elt2", LT2C, "LT2")):
    r0, r1 = lines[k]
    ch.series.append(mk_series(ws_cd, XC, YC, r0, r1, lab, colr, width=2))
DG.add_chart(ch, "A4")
ch = base_chart("Log-Log-Darstellung (Beaver)", "Belastung (log)", "Laktat [mmol/L] (log)", 16, 9)
ch.x_axis.scaling.logBase = 10
ch.y_axis.scaling.logBase = 10
ch.series.append(mk_series(ws_c, 22, 23, C0, C1, "Messwerte", INK, marker=True))
ch.series.append(mk_series(ws_cd, XC, YC, LLR, LLR + 1, "Regression unten", "E5482A", width=1.75))
ch.series.append(mk_series(ws_cd, XC, YC, LLR + 2, LLR + 3, "Regression oben", "E5482A", width=1.75))
DG.add_chart(ch, "J4")
DG["A24"] = '=IF(k_llok,"Log-Log-Breakpoint bei "&FIXED(k_llx,k_xd,TRUE)&" "&k_unit&" (Segmente "&k_llk&" + "&(k_n-k_llk)&" Punkte)","Log-Log-Breakpoint nicht bestimmbar")'
DG["A24"].font = f_base
DG["A25"] = f'=IF(k_ok,"LEmin bei "&FIXED(INDEX(VCalc!$C${LEM["x"]}:$F${LEM["x"]},k_model),k_xd,TRUE)&" "&k_unit&" ("&FIXED(INDEX(VCalc!$C${LEM["la"]}:$F${LEM["la"]},k_model),2,TRUE)&" mmol/L)","")'
DG["A25"].font = f_base
put(DG, "A27", "Glukose zu Belastung", f_h2)
DG["A28"] = ('=IF(AND(k_gluOn,k_ng>=2),k_ng&" Werte · erste → letzte Stufe "&IF(INDEX(Calc!$AL$3:$AL$22,k_ng)-Calc!$AL$3>=0,"+","")'
             '&FIXED(INDEX(Calc!$AL$3:$AL$22,k_ng)-Calc!$AL$3,k_gd,TRUE)&" "&k_gluU&" · Minimum "&FIXED(MIN(OFFSET(Calc!$AL$3,0,0,k_ng,1)),k_gd,TRUE)'
             '&" bei "&FIXED(INDEX(Calc!$AK$3:$AK$22,MATCH(MIN(OFFSET(Calc!$AL$3,0,0,k_ng,1)),OFFSET(Calc!$AL$3,0,0,k_ng,1),0)),k_xd,TRUE)&" "&k_unit,'
             '"Glukose-Auswertung: im Blatt „Eingabe“ „Glukose je Stufe erfassen = Ja“ wählen und mind. 2 Werte eintragen.")')
DG["A28"].font = f_base
ch = base_chart("Glukose zu Belastung", "Belastung", "Glukose", 16, 8)
ch.series.append(mk_series(ws_cd, 7, 8, 3, 203, "Interpolation (formtreu)", "8A5CC2", width=2))
ch.series.append(mk_series(ws_c, 44, 45, C0, C1, "Glukose gemessen", "8A5CC2", marker=True))
ch.series.append(mk_series(ws_cd, XC, YC, RG, RG + 1, "Ruheglukose", MUTED, dash="dash", width=1.2))
DG.add_chart(ch, "A30")
DG.column_dimensions["A"].width = 12

# ====================================================================== Bericht (druckfertig)
B = ws_rep
for col, w in zip("ABCDEFGH", (22, 16, 12, 12, 12, 12, 12, 12)):
    B.column_dimensions[col].width = w
B["A1"] = '=IF(Eingabe!$B$12<>"",Eingabe!$B$12,"")'
B["A1"].font = f_bold
B["A2"] = '=IF(Eingabe!$B$13<>"",Eingabe!$B$13,"")'
B["A2"].font = f_muted
put(B, "H1", "Laktatdiagnostik", Font(name=FONT, size=14, bold=True), align=Alignment(horizontal="right"))
put(B, "A3", "Ergebnisse der Leistungsdiagnostik", Font(name=FONT, size=13, color="FFFFFF"),
    PatternFill("solid", fgColor="BEC3CA"), align=Alignment(horizontal="center"))
B.merge_cells("A3:H3")
info = [("Name", "=Eingabe!$B$4", "Testleitung", "=Eingabe!$B$10"),
        ("Geburtsdatum", '=IF(ISNUMBER(Eingabe!$B$5),TEXT(Eingabe!$B$5,"DD.MM.YYYY")&IF(ISNUMBER(k_age),"  ("&FIXED(k_age,1,TRUE)&" Jahre)",""),Eingabe!$B$5&"")',
         "Untersuchung", '=IF(k_run,"Laufbandergometrie","Fahrradergometrie")'),
        ("Größe", '=IF(ISNUMBER(Eingabe!$B$7),Eingabe!$B$7&" cm","–")', "Stufen",
         '=IF(ISNUMBER(k_dur),FIXED(k_dur,1,TRUE),"–")&" min, "&IF(Eingabe!$E$5="","–",Eingabe!$E$5)'),
        ("Masse", '=IF(ISNUMBER(Eingabe!$B$8),Eingabe!$B$8&" kg","–")', "Probe", '=Eingabe!$E$6&", "&Eingabe!$E$7'),
        ("Testdatum", '=IF(ISNUMBER(Eingabe!$B$6),TEXT(Eingabe!$B$6,"DD.MM.YYYY"),Eingabe!$B$6&"")', "Analysator",
         '=IF(Eingabe!$E$8="","–",Eingabe!$E$8)'),
        ("Bemerkung", '=Eingabe!$B$11&""', "", ""),
        ("Ruhewerte", '="Laktat "&IF(ISNUMBER(k_rest),FIXED(k_rest,2,TRUE),"–")&" mmol/L     Glukose "&IF(ISNUMBER(Eingabe!$E$10),FIXED(Eingabe!$E$10,k_gd,TRUE),"–")&" "&k_gluU', "", "")]
for i, (a, b, c_, d) in enumerate(info):
    rr = 5 + i
    put(B, f"A{rr}", a, f_bold)
    B[f"B{rr}"] = b
    B[f"B{rr}"].font = f_base
    put(B, f"E{rr}", c_ or None, f_bold)
    if d:
        B[f"F{rr}"] = d
        B[f"F{rr}"].font = f_base
B["A11"].fill = fill_soft
B.merge_cells("B10:H10")
B.merge_cells("B11:H11")


def rep_h2(r, text):
    put(B, f"A{r}", text, Font(name=FONT, size=11, bold=True))
    for cc in "ABCDEFGH":
        B[f"{cc}{r}"].border = Border(bottom=Side(style="medium", color=INK))


rep_h2(13, "Gemessene Werte")
mh = ["Zeit [min]", '="Belastung ["&k_unit&"]"', "Pace [min/km]", "Laktat [mmol/L]", "HF [1/min]", '=IF(k_gluOn,"Glukose ["&k_gluU&"]","")']
for j, h in enumerate(mh):
    c = B.cell(14, 1 + j, h)
    c.font = Font(name=FONT, size=9, bold=True, color=MUTED)
    c.alignment = Alignment(horizontal="left" if j == 0 else "right", wrap_text=True)
for i in range(MAXST):
    rr = 15 + i
    k = i + 1
    src = f"Calc!$H${C0+i}"
    B[f"A{rr}"] = f'=IF({k}<=k_n,INDEX(Eingabe!$B${ST0}:$B${ST1},{src})&"","")'
    B[f"B{rr}"] = f'=IF({k}<=k_n,ROUND(Calc!$G${C0+i},k_xd),"")'
    B[f"C{rr}"] = f'=IF(AND(k_run,{k}<=k_n),{pace(f"Calc!$G${C0+i}")},"")'
    B[f"D{rr}"] = f'=IF({k}<=k_n,Calc!$I${C0+i},"")'
    B[f"E{rr}"] = f'=IF({k}<=k_n,IF(ISNUMBER(Calc!$J${C0+i}),Calc!$J${C0+i},"–"),"")'
    B[f"F{rr}"] = f'=IF(AND(k_gluOn,{k}<=k_n),IF(ISNUMBER(INDEX(Eingabe!$F${ST0}:$F${ST1},{src})),INDEX(Eingabe!$F${ST0}:$F${ST1},{src}),"–"),"")'
    for cc in "ABCDEF":
        B[f"{cc}{rr}"].font = Font(name=FONT, size=9)
        B[f"{cc}{rr}"].alignment = Alignment(horizontal="left" if cc == "A" else "right")
    B[f"D{rr}"].number_format = "0.00"
B.conditional_formatting.add(f"A15:F{14+MAXST}", FormulaRule(formula=['AND($B15<>"",MOD(ROW(),2)=0)'], fill=PatternFill("solid", fgColor="F7F8FA")))
TR = 15 + MAXST + 1
rep_h2(TR, "Schwellen und Maximalwerte")
th = ["", "Verfahren", '="Belastung ["&k_unit&"]"', "Pace [min/km]", "Laktat [mmol/L]", "HF [1/min]", "HF [%]"]
for j, h in enumerate(th):
    c = B.cell(TR + 1, 1 + j, h)
    c.font = Font(name=FONT, size=9, bold=True, color=MUTED)
    c.alignment = Alignment(horizontal="left" if j < 2 else "right", wrap_text=True)
x4 = f'INDEX(Calc!$D${LT0}:$D${LT1r},5)'
trows = [
    ("LT1", '=Eingabe!$I$5&IF(k_man1," (manuell)","")', "k_x1", "k_la1"),
    ("LT2", '=Eingabe!$I$6&IF(k_man2," (manuell)","")', "k_x2", "k_la2"),
    ("4 mmol/L", '="Vergleichspunkt"', f'IF(INDEX(Calc!$E${LT0}:$E${LT1r},5),{x4},"")', "4"),
    ("Maximum", '="letzte Stufe"', "k_hi", "k_ylast"),
]
for i, (lab, meth, xe, la) in enumerate(trows):
    rr = TR + 2 + i
    put(B, f"A{rr}", lab, Font(name=FONT, size=10, bold=True, color=[LT1C, LT2C, INK, INK][i]))
    B[f"B{rr}"] = f"=IF(k_ok,{meth[1:]},\"\")"
    hrv = hrat(xe) if i < 3 else "k_hrlast"
    B[f"C{rr}"] = f'=IF(AND(k_ok,ISNUMBER({xe})),ROUND({xe},k_xd),"–")'
    B[f"D{rr}"] = f'=IF(AND(k_ok,k_run,ISNUMBER({xe})),{pace(xe)},"–")'
    B[f"E{rr}"] = f'=IF(AND(k_ok,ISNUMBER({xe})),ROUND({la},2),"–")'
    B[f"F{rr}"] = f'=IF(AND(k_ok,ISNUMBER({xe})),IFERROR(ROUND({hrv},0),"–"),"–")'
    B[f"G{rr}"] = f'=IF(AND(ISNUMBER(F{rr}),ISNUMBER(k_hrmax)),ROUND(F{rr}/k_hrmax*100,1),"–")'
    for cc in "BCDEFG":
        B[f"{cc}{rr}"].font = Font(name=FONT, size=10)
        B[f"{cc}{rr}"].alignment = Alignment(horizontal="left" if cc == "B" else "right")
        B[f"{cc}{rr}"].border = b_bottom
    B[f"A{rr}"].border = b_bottom
ZR = TR + 7
rep_h2(ZR, "Trainingsbereiche")
zh = ["Trainingsbereich", "", "HF von", "HF bis", '="von ["&k_unit&"]"', '="bis ["&k_unit&"]"', "Pace von", "Pace bis"]
for j, h in enumerate(zh):
    c = B.cell(ZR + 1, 1 + j, h)
    c.font = Font(name=FONT, size=9, bold=True, color=MUTED)
    c.alignment = Alignment(horizontal="left" if j < 2 else "right")
for i in range(5):
    rr = ZR + 2 + i
    orow = ZT + 2 + i
    B[f"A{rr}"] = f'=Auswertung!$A${orow}&IF(k_zm{i}," *","")'
    B.merge_cells(f"A{rr}:B{rr}")
    for j, cc in enumerate("CDEFGH"):
        B[f"{cc}{rr}"] = f"=Auswertung!{'BCDEFG'[j]}{orow}"
        B[f"{cc}{rr}"].alignment = Alignment(horizontal="right")
    for cc in "ABCDEFGH":
        B[f"{cc}{rr}"].fill = PatternFill("solid", fgColor=ZC[i])
        B[f"{cc}{rr}"].font = Font(name=FONT, size=10, bold=(cc == "A"))
B[f"A{ZR+7}"] = '=IF(OR(k_zm0,k_zm1,k_zm2,k_zm3,k_zm4),"* Bereichsgrenze manuell angepasst","")'
B[f"A{ZR+7}"].font = f_muted
# Diagramme (Seite 2)
CHR = ZR + 9
B.row_breaks.append(__import__("openpyxl").worksheet.pagebreak.Break(id=CHR - 1))
rep_h2(CHR, "Laktat zu Belastung")
B.add_chart(la_chart(""), f"A{CHR+2}")
B[f"A{CHR+22}"] = '=IF(k_ok,"Modell: "&INDEX(Listen!$F$2:$F$5,k_model)&", R² "&FIXED(N(k_r2),4,TRUE)&", RMSE "&FIXED(k_rmse,3,TRUE)&" mmol/L"&IF(k_gok," · rot: Hilfsgeraden (RMSE "&FIXED(k_grmse,3,TRUE)&" mmol/L)",""),"")'
B[f"A{CHR+22}"].font = f_muted
rep_h2(CHR + 24, "Herzfrequenz zu Belastung")
B.add_chart(hr_chart(""), f"A{CHR+26}")
for ch in B._charts:
    ch.width, ch.height = 17.5, 10
# Seite 3: Tabelle, Erläuterungen, Methodik, Hinweise
P3 = CHR + 47
B.row_breaks.append(__import__("openpyxl").worksheet.pagebreak.Break(id=P3 - 1))
rep_h2(P3, "Tabellarische Aufstellung")
for j, h in enumerate(["Laktat [mmol/L]", '="Belastung ["&k_unit&"]"', "Pace [min/km]", "Belastung [m/s]", "HF [1/min]"]):
    c = B.cell(P3 + 1, 1 + j, h)
    c.font = Font(name=FONT, size=9, bold=True, color=MUTED)
    c.alignment = Alignment(horizontal="left" if j == 0 else "right")
for i in range(20):
    rr = P3 + 2 + i
    for j, cc in enumerate("ABCDE"):
        B[f"{cc}{rr}"] = f'=Auswertung!{cc}{LVT+2+i}&""' if cc in "C" else f"=Auswertung!{cc}{LVT+2+i}"
        B[f"{cc}{rr}"].font = Font(name=FONT, size=9)
        B[f"{cc}{rr}"].alignment = Alignment(horizontal="left" if cc == "A" else "right")
    B[f"A{rr}"].number_format = "0.0"
EX = P3 + 23
rep_h2(EX, "Erläuterung der Trainingsbereiche")
expl = [("A0 Regeneration/Kompensation", "Unterhalb der Trainingsreizschwelle. Dient der Erholung nach intensiven Einheiten: kurz, locker, gleichmäßig. Beispiel: 15–20 min lockeres Laufen oder 30 min lockeres Radfahren."),
        ("A1 Extensives Ausdauertraining 1", "Aufbau der Grundlagenausdauer mit langen, ruhigen Einheiten bei geringer Intensität (Dauermethode, konstantes oder leicht wechselndes Tempo). Beispiel: 60–120 min Dauerlauf."),
        ("A2 Extensives Ausdauertraining 2", "Ausbau der Ausdauerbasis bei mittlerer Intensität, kürzer als A1. Dauermethode oder Wechseltempo zwischen unterer und oberer Bereichsgrenze. Beispiel: 45–60 min Dauerlauf."),
        ("A3 Intensives Ausdauertraining", "Training im Bereich um die obere Schwelle, mittlere bis hohe Intensität. Dauer- oder Intervallmethode. Beispiel: 30–60 min zügiger Dauerlauf oder Intervalle mit lohnender Pause."),
        ("A4 Entwicklungsbereich", "Oberhalb der Schwelle, hohe Intensität, kurze Belastungen. Vor allem Intervallmethode mit ausreichenden Pausen. Beispiel: 10–30 min Gesamtbelastung in Intervallen.")]
for i, (k, v) in enumerate(expl):
    rr = EX + 1 + i
    put(B, f"A{rr}", k, Font(name=FONT, size=9, bold=True), align=Alignment(vertical="top", wrap_text=True))
    put(B, f"B{rr}", v, Font(name=FONT, size=9), align=Alignment(vertical="top", wrap_text=True))
    B.merge_cells(f"B{rr}:H{rr}")
    B.row_dimensions[rr].height = 26
ME = EX + 7
rep_h2(ME, "Methodik")
for i in range(len(proto)):
    rr = ME + 1 + i
    B[f"A{rr}"] = f'=IF(Calc!$B${PR0+i}<>"",Calc!$A${PR0+i},"")'
    B[f"A{rr}"].font = Font(name=FONT, size=9, bold=True)
    B[f"A{rr}"].alignment = Alignment(vertical="top")
    B[f"B{rr}"] = f"=Calc!$B${PR0+i}"
    B[f"B{rr}"].font = Font(name=FONT, size=9)
    B[f"B{rr}"].alignment = Alignment(vertical="top", wrap_text=True)
    B.merge_cells(f"B{rr}:H{rr}")
    B.row_dimensions[rr].height = {"Protokoll": 14, "Analytik": 14, "LTP-Konzept": 50, "Hilfsgeraden": 62,
                                   "Dmax-Verfahren": 38}.get(proto[i][0], 26)
HN = ME + len(proto) + 2
rep_h2(HN, "Hinweise zur Auswertung")
for i in range(len(notes)):
    rr = HN + 1 + i
    B[f"A{rr}"] = f'=IF(Auswertung!$A${NH+1+i}<>"","– "&Auswertung!$A${NH+1+i},"")'
    B[f"A{rr}"].font = Font(name=FONT, size=9)
B.print_area = f"A1:H{HN+len(notes)}"
B.page_setup.paperSize = B.PAPERSIZE_A4
B.page_setup.orientation = "portrait"
B.page_setup.fitToWidth = 1
B.page_setup.fitToHeight = 0
B.sheet_properties.pageSetUpPr.fitToPage = True
B.print_options.horizontalCentered = True
B.page_margins.left = B.page_margins.right = 0.5
B.oddFooter.left.text = "&8" + "Laktatdiagnostik"
B.oddFooter.center.text = "&8Erstellt &D"
B.oddFooter.right.text = "&8Seite &P von &N"

# ====================================================================== Beispiele
X2 = ws_ex
put(X2, "A1", "Beispieldaten aus der Webapp", f_head)
put(X2, "A2", "Zum Ausprobieren: Werte in die gelben Felder des Blatts „Eingabe“ kopieren (Messwerte nach Eingabe!B17, Stammdaten einzeln).", f_muted)
rr = 4
for title, ex in EXAMPLES.items():
    put(X2, f"A{rr}", title, f_h2)
    rr += 1
    for k, v in ex["meta"].items():
        put(X2, f"A{rr}", {"name": "Name", "birth": "Geburtsdatum", "date": "Testdatum", "height": "Größe [cm]",
                           "mass": "Masse [kg]", "tester": "Testleitung", "stageInc": "Stufenhöhe",
                           "restLa": "Ruhelaktat [mmol/L]", "restGlu": "Ruheglukose [mg/dL]"}[k])
        put(X2, f"B{rr}", v)
        rr += 1
    for j, h in enumerate(["Zeit", "Belastung [km/h]", "Laktat [mmol/L]", "HF [1/min]"]):
        put(X2, f"{CL(2+j)}{rr}", h, f_bold)
    rr += 1
    for s in ex["stages"]:
        for j, v in enumerate(s):
            put(X2, f"{CL(2+j)}{rr}", v)
        rr += 1
    rr += 2
X2.column_dimensions["A"].width = 24
for cc in "BCDE":
    X2.column_dimensions[cc].width = 16

# ====================================================================== Anleitung
A = ws_help
A.column_dimensions["A"].width = 120
lines_help = [
    ("Laktatdiagnostik – Excel-Version", f_head),
    ("Nachbau der Webapp „Laktatdiagnostik – Auswertung“ (Stufentest mit LT1/LT2, Modellvergleich, Trainingsbereichen). Alle Ergebnisse sind Formeln ohne Makros.", f_base),
    ("", f_base),
    ("So wird die Mappe benutzt", f_h2),
    ("1. Blatt „Eingabe“: Stammdaten, Protokoll und die Messwerte je Stufe (Belastung, Laktat, HF, optional Glukose) in die gelben Felder eintragen.", f_base),
    ("2. Rechts im Blatt „Eingabe“ Kurvenmodell, LT1-/LT2-Verfahren und die übrigen Optionen aus den Auswahllisten wählen.", f_base),
    ("3. Blatt „Auswertung“: Schwellen, Hinweise, Trainingsbereiche, Laktatstufen-Tabelle sowie Laktat- und HF-Diagramm.", f_base),
    ("4. Blatt „Verfahren“: alle 21 Schwellenverfahren und der Modellvergleich (Poly 2/3/4, Exponentiell).", f_base),
    ("5. Blatt „Diagramme“: Laktatäquivalent, Log-Log-Darstellung und Glukoseverlauf.", f_base),
    ("6. Blatt „Bericht“: druckfertiger A4-Bericht (Datei › Drucken oder Als PDF speichern).", f_base),
    ("7. LT1/LT2 manuell verschieben: Wert in Eingabe!I15 bzw. I16 eintragen (leer = Algorithmus). Bereichsgrenzen manuell: Eingabe!I26–I30.", f_base),
    ("8. Tests speichern: Datei unter neuem Namen sichern (ersetzt „Test speichern“ und JSON-Export der Webapp). Beispieldaten im Blatt „Beispiele“.", f_base),
    ("", f_base),
    ("Unterschiede zur Webapp", f_h2),
    ("• Nullstellen, Dmax, Laktatäquivalent-Minimum und Laktatminimum werden auf einem Raster mit 1000 Schritten im Messbereich gesucht und interpoliert bzw. per Parabel verfeinert (Abweichung zur Webapp typischerweise < 0,01 km/h bzw. < 0,1 W).", f_base),
    ("• 3-Phasen-Regression (LTP1/LTP2 global) und Hofmann/Pokan-LTP: Rastersuche mit Verfeinerung wie in der Webapp; bei flachen Fehlerflächen können einzelne Knoten um wenige Hundertstel abweichen.", f_base),
    ("• Exponentialfit: Rastersuche über c mit zweifacher Verfeinerung statt Goldener-Schnitt-Suche.", f_base),
    ("• Diagramme: Trainingsbereiche als gestrichelte Grenzlinien statt Farbflächen; Linien werden nicht mit der Maus gezogen, sondern über Eingabe!I15/I16 bzw. I26–I30 gesetzt.", f_base),
    ("• Die Grenze A1|A2 kann manuell gesetzt werden (Eingabe!I27). Eine HF-Eingabe für Bereichsgrenzen gibt es nicht; bitte Belastungswerte eintragen.", f_base),
    ("• Gespeicherte Tests im Browser, JSON-Import/-Export und die Beispielauswahl entfallen; stattdessen einzelne Excel-Dateien und das Blatt „Beispiele“.", f_base),
    ("• Ausgeblendete Rechenblätter (Calc, Grid, ExpFit, LTP, LTP3, Guide, VCalc, ChartData, Listen) enthalten die Zwischenrechnungen. Nicht verändern.", f_base),
    ("", f_base),
    ("Farblegende", f_h2),
    ("Gelb hinterlegt / blaue Schrift = Eingabefeld. Schwarze Schrift = Formel.", f_base),
    ("", f_base),
    ("Fachlicher Hinweis", f_h2),
    ("Die Schwellen sind modell- und verfahrensabhängige Schätzungen. LT2 ist nicht per Definition MLSS oder Critical Speed/Power. Die Voreinstellungen der Trainingsbereiche sind Konvention.", f_base),
]
for i, (t, f) in enumerate(lines_help):
    c = put(A, f"A{1+i}", t, f)
    c.alignment = Alignment(wrap_text=True, vertical="top")

# ====================================================================== Vorbelegung: Beispiel Mayer
ex = EXAMPLES["Mayer Elisabeth · Laufband 22.09.2026"]
import datetime as _dt
E["B4"] = ex["meta"]["name"]
E["B5"] = _dt.datetime(2012, 7, 15)
E["B6"] = _dt.datetime(2026, 9, 22)
E["B7"] = ex["meta"]["height"]
E["B8"] = ex["meta"]["mass"]
E["B9"] = "Laufband (km/h)"
E["B10"] = ex["meta"]["tester"]
E["E4"] = 3
E["E5"] = ex["meta"]["stageInc"]
E["E6"] = "Kapillar Ohrläppchen"
E["E7"] = "Vollblut"
E["E9"] = ex["meta"]["restLa"]
E["E10"] = ex["meta"]["restGlu"]
for i, s in enumerate(ex["stages"]):
    rr = ST0 + i
    E[f"B{rr}"], E[f"C{rr}"], E[f"D{rr}"], E[f"E{rr}"] = s

wb.active = 1
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("gespeichert:", OUT)
